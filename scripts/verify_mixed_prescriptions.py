"""Check mixed public Korean product labels against an independent coefficient-table oracle."""
import argparse
from collections import defaultdict
import csv
import io
import json
import math
from pathlib import Path
import sys
from urllib.request import Request, urlopen
import uuid

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app import app
from openpyxl import load_workbook
from services.korean_products import CATALOG

# Freeze the agreed order here: a production reorder must fail this verification.
ORDER = ('ED95', 'MED', 'GARDNER', 'INADA', 'CMD', 'CMD_DIRECT', 'CMD_INDIRECT', 'DDD', 'WOODS', 'CPZ_FGA')
# INADA is a supplement-only table (blonanserin); it is neither a selectable base nor a pure-method column.
BASES = tuple(m for m in ORDER if m != 'INADA')
TARGET = {m: 'chlorpromazine' if m in ('WOODS', 'CPZ_FGA') else 'olanzapine' for m in ORDER}
SHORT = {'olanzapine': 'OLZ', 'risperidone': 'RIS', 'chlorpromazine': 'CPZ'}
with (ROOT/'lookup/master_lookup.csv').open(encoding='utf-8-sig') as stream:
    FACTORS = {(r['method_id'], r['source_drug'], r['target_drug']): float(r['factor']) for r in csv.DictReader(stream)}
PRODUCTS = {p['code']: p for p in CATALOG['products']}


def product(code, frequency='QD', daily=None, active=None, days=None, oral=None, depot=None, held=False):
    p = PRODUCTS[code]
    return dict(text=(p['name']+' '+frequency).strip(), drug=p['drug'],
                route='injection' if p['route']=='주사' else 'oral',
                daily=daily if daily is not None else p['dose_mg'],
                active=active, days=days, oral=oral, depot=depot, held=held,
                tabs=(daily/p['dose_mg']) if daily is not None and p['route']!='주사' else 1)


def regimens():
    olz = product('641908390')
    cloz = product('642706240', 'BID', daily=50)
    ris = product('646902470')
    pp1 = product('646901591', 'q1mo', active=100, days=30, oral=9, depot=2.5)
    pp3 = product('646902161', 'q3mo', active=263, days=90, oral=6, depot=2.5)
    pp6 = product('646902671', 'q6mo', active=700, days=180, oral=9, depot=2.5)
    ari = product('649900421', 'q4w', active=400, days=28, depot=13.3)
    hal = product('076100071', 'q4w', active=50, days=28, depot=3.3)
    amis = product('651903820')
    clop = dict(text='clopenthixol 60mg tablet QD', drug='clopenthixol', route='oral', daily=60, tabs=1)
    unknown = dict(text='unknownxyz 5mg', drug='unknownxyz', held=True, tabs=1)
    lithium = dict(text='lithium 600mg QD', drug='lithium', excluded=True, tabs=1)
    pal_code = next(p['code'] for p in CATALOG['products'] if p['drug']=='paliperidone' and p['route']=='내복' and p['dose_mg']==6)
    return [
        dict(id='M001', name='경구 두 성분과 월간 주사', rows=[olz, cloz, pp1]),
        dict(id='M002', name='여러 단계의 보완', rows=[olz, pp3, ari, clop]),
        dict(id='M003', name='반년 주사와 경구약', rows=[amis, ris, pp6]),
        dict(id='M004', name='같은 성분의 경구약과 주사', rows=[product(pal_code), pp1, cloz]),
        dict(id='M005', name='처방 순서를 뒤집어도 같은 합계', rows=[clop, ari, pp3, olz]),
        dict(id='M006', name='기준값이 모두 있는 경구 대조', rows=[olz, ris]),
        dict(id='M007', name='단기 주사 미환산', rows=[olz, pp1, product('657200921', '', held=True)]),
        dict(id='M008', name='잘못된 주사 간격', rows=[cloz, ris, product('649900421', 'q1w', held=True)]),
        dict(id='M009', name='미확인 약물이 남는 부분합', rows=[olz, pp3, unknown]),
        dict(id='M010', name='대상 외 약물의 명시적 0', rows=[lithium, pp1, ris]),
        dict(id='M011', name='같은 상품의 서로 다른 함량', rows=[product('642706240'), product('642706610'), ari]),
        dict(id='M012', name='두 지속형 주사와 경구약', rows=[amis, hal, pp3]),
    ]


def pure_value(row, method, target):
    if row.get('held'):
        return None
    if row.get('excluded'):
        return 0.0
    drug, dose = row['drug'], row.get('daily')
    if row['route']=='injection':
        if method=='DDD':
            drug, dose = 'chlorpromazine', row['active']/row['days']/row['depot']*300
        else:
            dose = row['oral']
    factor = FACTORS.get((method, drug, target))
    return dose*factor if dose is not None and factor is not None else None


def expected(row, base):
    if row.get('excluded'):
        return '', 0.0
    values = [(m, pure_value(row, m, TARGET[base])) for m in (base, *[m for m in ORDER if m!=base])]
    return next(((m, value) for m, value in values if value is not None), ('', None))


def same(actual, expected_value):
    return actual is None if expected_value is None else isinstance(actual, (float, int)) and math.isclose(actual, expected_value, rel_tol=1e-10, abs_tol=1e-9)


def check_case(case, prescription, base):
    expected_rows = [expected(r, base) for r in prescription['rows']]
    assert case['base']==base and case['supplement_order']==[m for m in ORDER if m!=base]
    assert len(case['details'])==len(expected_rows), (prescription['id'], base, 'split/merged medications')
    by_text = defaultdict(list)
    for detail in case['details']:
        by_text[detail['original']].append(detail)
    for row, (method, value) in zip(prescription['rows'], expected_rows):
        # Date exports place non-target medicines last; match by retained input.
        assert by_text[row['text']], (prescription['id'], 'original input not preserved', row['text'])
        actual = by_text[row['text']].pop(0)
        assert actual['actual']==method and same(actual['value'], value), (prescription['id'], base, row['text'], actual, (method,value))
        assert actual['supplemented']==bool(method and method!=base)
        for pure in BASES:
            assert same(actual[pure], pure_value(row, pure, TARGET[pure])), (prescription['id'], base, pure)
    values = [v for _, v in expected_rows if v is not None]
    missing = len(expected_rows)-len(values)
    amount = round(math.fsum(values),4) if values else None
    assert same(case['total'], None if missing else amount)
    assert same(case['partial'], amount if missing else None)
    assert case['missing']==missing
    fallbacks = [m for m in ORDER if m!=base and any(actual==m for actual,_ in expected_rows)]
    assert case['complement']==' → '.join(fallbacks)
    assert case['supplemented']==sum(bool(m and m!=base) for m,_ in expected_rows)


def table(wb, name):
    rows = list(wb[name].values)
    return [dict(zip(rows[0], row)) for row in rows[1:]]


def check_workbook(content, prescriptions, base, quick=False):
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames==['보완합계','보완근거','Results','MedicationResults','Review','AuditTrail','코드안내']
    summaries, details = table(wb,'보완합계'), table(wb,'보완근거')
    assert len(summaries)==len(prescriptions)
    assert len(details)==sum(len(p['rows']) for p in prescriptions)
    for p in prescriptions:
        pid = 'P001' if quick else p['id']
        summary, = [r for r in summaries if r['환자 ID']==pid]
        selected = [r for r in details if r['환자 ID']==pid]
        # Reuse the independent arithmetic checks for the spreadsheet representation.
        case = dict(base=summary['기본 방법'], total=summary['보완 총합 (mg/day)'], partial=summary['부분합 (mg/day)'],
                    complement=summary['보완 방법'] or '', missing=summary['미환산 수'], supplemented=summary['보완 약물 수'],
                    supplement_order=[m for m in ORDER if m!=base], details=[])
        for row in selected:
            case['details'].append(dict(original=row['원문 약물'], actual=row['적용 방법'] or '', value=row['환산값 (mg/day)'],
                supplemented=row['처리']=='보완', **{m:row[f'{m} ({SHORT[TARGET[m]]} mg/day)'] for m in BASES}))
        check_case(case,p,base)
    return wb


def upload_data(prescriptions, dated=False):
    stream = io.StringIO(newline='')
    writer = csv.writer(stream)
    writer.writerow(['HID','PRESCR_DATE','DRUG','TABS_PER_DAY'] if dated else ['patient_id','medication'])
    for p in prescriptions:
        for row in p['rows']:
            writer.writerow([p['id'],'2026-10-20',row['text'],row['tabs']] if dated else [p['id'],row['text']])
    return stream.getvalue().encode('utf-8-sig')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base'); parser.add_argument('--commit')
    parser.add_argument('--output',type=Path,required=True)
    args = parser.parse_args()
    client, prescriptions = app.test_client(), regimens()
    def request(path, data=None, kind=None):
        with urlopen(Request(args.base.rstrip('/')+path,data=data,headers={'Content-Type':kind} if kind else {}),timeout=90) as response:
            return response.read()
    def api(path, payload):
        return request(path,json.dumps(payload).encode(),'application/json') if args.base else client.post(path,json=payload).data
    release = json.loads(request('/version')) if args.base else client.get('/version').json
    if args.commit: assert release['commit']==args.commit
    checked, exports = [], 0
    for base in BASES:
        for p in prescriptions:
            payload = dict(text=', '.join(r['text'] for r in p['rows']),base_method=base)
            data = json.loads(api('/api/parse',payload))
            case, = data['cases']
            check_case(case,p,base)
            checked.append(dict(id=p['id'],base=base,total=case['total'],partial=case['partial'],
                                methods=[r['actual'] for r in case['details']],missing=case['missing']))
        p = prescriptions[1]
        content = api('/api/export',dict(text=', '.join(r['text'] for r in p['rows']),base_method=base))
        check_workbook(content,[p],base,quick=True); exports += 1
        for dated in (False,True):
            raw = upload_data(prescriptions,dated)
            local = client.post('/upload',data={'base_method':base,'file':(io.BytesIO(raw),'mixed.csv')})
            assert local.status_code==200,local.data[:500]
            content = local.data
            if args.base:
                boundary = 'ap-mixed-'+uuid.uuid4().hex
                body = (f'--{boundary}\r\nContent-Disposition: form-data; name="base_method"\r\n\r\n{base}\r\n'
                        f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="mixed.csv"\r\nContent-Type: text/csv\r\n\r\n').encode()+raw+f'\r\n--{boundary}--\r\n'.encode()
                content = request('/upload',body,'multipart/form-data; boundary='+boundary)
            actual = check_workbook(content,prescriptions,base)
            expected_book = load_workbook(io.BytesIO(local.data))
            assert all(list(actual[s].values)==list(expected_book[s].values) for s in actual.sheetnames)
            exports += 1
        print(f'{base}: {len(prescriptions)} mixed prescriptions and 3 workbooks passed',flush=True)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    report = dict(release=release,prescriptions=len(prescriptions),bases=len(BASES),checks=len(checked),
                  workbooks=exports,passed=True,source=CATALOG['source'],synthetic_conditions_only=True,results=checked)
    args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')


if __name__=='__main__': main()
