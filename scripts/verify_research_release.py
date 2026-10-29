"""Compare synthetic local/live API and every Excel worksheet after exact version match."""
import argparse
import io
import json
from pathlib import Path
import sys
from urllib.request import Request, urlopen
import uuid

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app import app
from services.release import VERSION
from openpyxl import load_workbook

CASES = [
    'risperidone 2mg QD, olanzapine 5mg QD',
    'risperidone 2000㎍ QD', 'risperidone 2mg/mL QD',
    'quetiapine take 1/2 tablet (50 mg) every day',
    'risperidone 2mg q2-3hr', 'risperidone 2mg QD, unknownzapine 5mg',
    'Sustenna 156mg monthly', 'Aristada 441mg monthly',
    'Rykindo 37.5mg q2w', 'Longavo 50mg',
    'lithium 600mg', 'paliperidone 6mg QD',
]
SOURCES = {
    'quick': 'patient_id,medication\n001,risperidone 2mg QD\n001,paliperidone 6mg QD\n002,unknownzapine 5mg\n003,lithium 600mg\n004,Sustenna 156mg monthly\n',
    'dated': 'HID,PRESCR_DATE,DRUG,TABS_PER_DAY,COMMERCIAL_DRUG\n001,2026-10-19,Risperidone 2mg tab,1,\n002,2026-10-19,Paliperidone 6mg tab,1,\n003,2026-10-19,unknownzapine 5mg tab,1,\n004,2026-10-19,risperidone 75mg,,Longavo 75mg monthly\n',
    'korean-structured': 'patient_id,drug,dose,unit,frequency\n001,클자핀정25밀리그램(클로자핀)_(25mg/1정),25,mg,BID\n002,아빌리파이메인테나주사400밀리그램,,,q4w\n003,클자핀정25밀리그램,50,mg,QD\n',
}


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base', default='https://ap-equivalence-calculator.onrender.com')
    parser.add_argument('--commit', required=True)
    parser.add_argument('--output', type=Path, default=ROOT/'output/research_live_verification.json')
    args=parser.parse_args()
    def request(path, data=None, content_type=None):
        req=Request(args.base.rstrip('/')+path, data=data, headers={'Content-Type':content_type} if content_type else {})
        with urlopen(req, timeout=40) as response:
            return response.read()
    release=json.loads(request('/version'))
    assert release['commit']==args.commit and release['version']==VERSION, release
    assert json.loads(request('/health'))['status']=='ok'
    local=app.test_client()
    assert json.loads(request('/api/methods'))==local.get('/api/methods').json
    assert request('/evidence.csv').replace(b'\r\n',b'\n')==(ROOT/'lookup/master_lookup.csv').read_bytes().replace(b'\r\n',b'\n')
    assert '환산 근거와 적용 한계' in request('/methods').decode()
    assert b'include_sensitivity' not in request('/')
    assert b'include-sensitivity' not in request('/longitudinal')
    assert 'sensitivity_version' not in release
    for i,text in enumerate(CASES,1):
        from services.result_summary import METHOD_ORDER
        payload={'text':text,'base_method':METHOD_ORDER[(i-1)%len(METHOD_ORDER)]}
        expected=local.post('/api/parse',json=payload).json
        actual=json.loads(request('/api/parse',json.dumps(payload).encode(),'application/json'))
        assert actual['totals']==expected['totals'], i
        assert actual['cases']==expected['cases'], i
        assert len(actual['items'])==len(expected['items'])
        for a,b in zip(actual['items'],expected['items']):
            for key in ('drug','status','needs_review','warning','calculation_note','dose_mg','daily_dose_mg',
                        'active_moiety_mg','interval_days','conversions','factor_evidence'):
                assert a.get(key)==b.get(key), (i,key)
        print(f'API {i}/{len(CASES)} matched',flush=True)
    checked=[]
    def compare_excel(expected, actual, label):
        a,b=load_workbook(io.BytesIO(expected)),load_workbook(io.BytesIO(actual))
        assert a.sheetnames==b.sheetnames==['보완합계','보완근거','Results','MedicationResults','Review','AuditTrail']
        for name in a.sheetnames:
            assert list(a[name].values)==list(b[name].values), (label,name)
        checked.append(dict(case=label,sheets=a.sheetnames))
        print(f'Excel {label}: {len(a.sheetnames)} worksheets matched',flush=True)
    for name,text in SOURCES.items():
        selected='ED95' if name=='dated' else 'MED'
        raw=text.encode()
        expected=local.post('/upload',data={'file':(io.BytesIO(raw),'synthetic.csv'),'base_method':selected})
        assert expected.status_code==200, expected.data[:1000]
        boundary='ap-synthetic-'+uuid.uuid4().hex
        body=(f'--{boundary}\r\nContent-Disposition: form-data; name="base_method"\r\n\r\n{selected}\r\n--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="synthetic.csv"\r\nContent-Type: text/csv\r\n\r\n').encode()+raw+f'\r\n--{boundary}--\r\n'.encode()
        actual=request('/upload',body,'multipart/form-data; boundary='+boundary)
        compare_excel(expected.data,actual,name)
    payload=dict(text='risperidone 2mg QD, paliperidone 6mg QD',base_method='ED95')
    compare_excel(local.post('/api/export',json=payload).data,
                  request('/api/export',json.dumps(payload).encode(),'application/json'),'quick-api')
    report=dict(commit=args.commit,version=VERSION,api_cases=len(CASES),workbooks=checked,
                all_worksheets_match=True,health='ok',evidence_download_matches=True,
                synthetic_only=True)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False),flush=True)


if __name__=='__main__':
    main()
