"""Check frozen public product descriptors locally or on an exact live release."""
import argparse
from concurrent.futures import ThreadPoolExecutor
import csv
import io
import json
from pathlib import Path
import sys
from urllib.request import Request, urlopen
import uuid

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from app import app
from openpyxl import load_workbook

FIXTURE=ROOT/'data/evaluation/external_products.json'
SHEETS=['보완합계','보완근거','Results','MedicationResults','Review','AuditTrail']


def check_items(case, items):
    failures=[]
    if case['kind']=='no_antipsychotic_conversion':
        if any(c.get('value') not in (None,0) for i in items for c in i.get('conversions',[])):
            failures.append('non-antipsychotic descriptor received a positive equivalent')
        return failures
    if len(items)!=1:
        failures.append(f'expected one product, got {len(items)} frames')
    if case['expected_drug'] and not any(i.get('drug')==case['expected_drug'] for i in items):
        failures.append('ingredient not preserved')
    if case['kind']=='hold':
        if any(i.get('daily_dose_mg') is not None or any(c.get('value') is not None for c in i.get('conversions',[])) for i in items):
            failures.append('unsupported or incomplete product was calculated')
        if not all(i.get('needs_review') for i in items):
            failures.append('review flag missing')
    elif len(items)==1:
        if items[0].get('daily_dose_mg')!=case['daily_mg']:
            failures.append('one-unit synthetic daily mg mismatch')
        if items[0].get('needs_review'):
            failures.append('unnecessary input review')
    return failures


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base')
    parser.add_argument('--commit')
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--report-only',action='store_true')
    args=parser.parse_args()
    dataset=json.loads(FIXTURE.read_text(encoding='utf-8'))
    cases=dataset['cases']
    client=app.test_client()
    def request(path,data=None,content_type=None):
        req=Request(args.base.rstrip('/')+path,data=data,headers={'Content-Type':content_type} if content_type else {})
        with urlopen(req,timeout=60) as response:
            return response.read()
    release=json.loads(request('/version')) if args.base else client.get('/version').json
    if args.commit:
        assert release['commit']==args.commit,release
    def check(case):
        payload={'text':case['text']}
        response=json.loads(request('/api/parse',json.dumps(payload).encode(),'application/json')) if args.base else client.post('/api/parse',json=payload).json
        return dict(id=case['id'],items=response['items'],errors=check_items(case,response['items']))
    if args.base:
        with ThreadPoolExecutor(max_workers=2) as pool:
            checked=list(pool.map(check,cases))
    else:
        checked=[check(c) for c in cases]
    args.output.parent.mkdir(parents=True,exist_ok=True)
    buf=io.StringIO(newline='')
    writer=csv.writer(buf)
    writer.writerow(['patient_id','medication'])
    writer.writerows((c['id'],c['text']) for c in cases)
    raw=buf.getvalue().encode('utf-8-sig')
    (args.output.parent/'external_products_input.csv').write_bytes(raw)
    if args.base:
        boundary='public-products-'+uuid.uuid4().hex
        body=(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="external_products_input.csv"\r\nContent-Type: text/csv\r\n\r\n').encode()+raw+f'\r\n--{boundary}--\r\n'.encode()
        content=request('/upload',body,'multipart/form-data; boundary='+boundary)
    else:
        response=client.post('/upload',data={'file':(io.BytesIO(raw),'external_products_input.csv')})
        assert response.status_code==200,response.data[:1000]
        content=response.data
    (args.output.parent/'external_products_result.xlsx').write_bytes(content)
    wb=load_workbook(io.BytesIO(content))
    assert wb.sheetnames==SHEETS
    values=list(wb['AuditTrail'].values)
    audit=[dict(zip(values[0],r)) for r in values[1:]]
    excel_errors=[]
    for case,checked_case in zip(cases,checked):
        rows=[r for r in audit if r['patient']==case['id']]
        items=checked_case['items']
        if len(rows)!=len(items):
            excel_errors.append(dict(id=case['id'],error='API/Excel row count mismatch'))
            continue
        for row,item in zip(rows,items):
            for column,field in [('parsed','drug'),('daily_dose_mg','daily_dose_mg'),('needs_review','needs_review'),('route','route')]:
                if row[column] != item.get(field):
                    excel_errors.append(dict(id=case['id'],error=f'API/Excel {column} mismatch'))
    failures=[dict(id=r['id'],errors=r['errors']) for r in checked if r['errors']]
    summary=dict(cases=len(cases),api_passed=len(cases)-len(failures),api_failures=failures,
                 excel_api_mismatches=excel_errors,sheets=wb.sheetnames,release=release,
                 source_urls=dataset['source_urls'],clinical_validation=False,
                 public_product_data_and_synthetic_instructions_only=True)
    args.output.write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(dict(cases=len(cases),passed=summary['api_passed'],api_failures=len(failures),excel_mismatches=len(excel_errors))),flush=True)
    if not args.report_only:
        assert not failures and not excel_errors, 'See report for failures'


if __name__=='__main__':
    main()
