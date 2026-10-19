"""Verify the public MIMIC-IV demo's medication patterns, without patient fields.

Source: https://physionet.org/content/mimic-iv-demo/2.2/ (ODbL 1.0).
Download hosp/prescriptions.csv.gz separately. Neither source rows nor patient
identifiers are committed. Only six medication fields are read into test cases.
The adapter uses prescribed mg, not product strength or dispensed tablet count.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor
import csv
import gzip
import hashlib
import io
import json
from pathlib import Path
import re
import sys
from urllib.request import Request, urlopen
import uuid

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app import app
from openpyxl import load_workbook

SOURCE_URL = 'https://physionet.org/files/mimic-iv-demo/2.2/hosp/prescriptions.csv.gz'
SOURCE_SHA256 = '33c392ba5b9299b08eca0a61911ba106f0aebdba26ed31b856bb9ffd49fe3654'
FIELDS = ['drug', 'prod_strength', 'dose_val_rx', 'dose_unit_rx', 'doses_per_24_hrs', 'route']
NAMES = ('haloperidol', 'quetiapine', 'risperidone', 'olanzapine', 'aripiprazole',
         'ziprasidone', 'clozapine', 'chlorpromazine', 'perphenazine', 'fluphenazine',
         'lurasidone', 'paliperidone')
SHEETS = ['보완합계', '보완근거', 'Results', 'MedicationResults', 'Review', 'AuditTrail']


def load_cases(path):
    assert hashlib.sha256(path.read_bytes()).hexdigest() == SOURCE_SHA256, 'Unexpected source version'
    unique, selected, total = set(), 0, 0
    with gzip.open(path, 'rt', encoding='utf-8') as source:
        for row in csv.DictReader(source):
            total += 1
            if any(name in row['drug'].casefold() for name in NAMES):
                selected += 1
                unique.add(tuple(row[k] for k in FIELDS))
    cases = []
    for index, values in enumerate(sorted(unique), 1):
        row = dict(zip(FIELDS, values))
        freq = row['doses_per_24_hrs']
        text = f"{row['drug']} {row['dose_val_rx']}{row['dose_unit_rx']} {row['route']}"
        if freq:
            text += f' {freq} times a day'
        cases.append(dict(id=f'PUBLIC{index:03}', text=text, source=row))
    return cases, dict(source_rows=total, selected_rows=selected, unique_patterns=len(cases))


def errors(case, items):
    row = case['source']
    if len(items) != 1:
        return [f'one medication split into {len(items)} items']
    item = items[0]
    result = []
    drug = next(name for name in NAMES if name in row['drug'].casefold())
    if item.get('drug') != drug:
        result.append('incorrect drug')
    injected = row['route'] in {'IV', 'IM'}
    if item.get('route') != ('injection' if injected else 'oral'):
        result.append('incorrect route')
    numeric = re.fullmatch(r'\d+(?:\.\d+)?', row['dose_val_rx']) is not None
    freq = row['doses_per_24_hrs']
    should_hold = injected or not numeric or freq == '0'
    if should_hold:
        if item.get('daily_dose_mg') is not None or any(c['value'] is not None for c in item.get('conversions', [])):
            result.append('unresolved dose or injection was calculated')
        if not item.get('needs_review'):
            result.append('missing review flag')
    else:
        assert row['dose_unit_rx'] == 'mg', 'Update independent unit expectations before adding new units'
        expected = float(row['dose_val_rx']) * (float(freq) if freq else 1)
        if item.get('daily_dose_mg') != expected:
            result.append(f'daily dose should be {expected}')
        if item.get('needs_review') != (not bool(freq)):
            result.append('incorrect frequency review flag')
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--base', help='Optional live URL; otherwise use local Flask client')
    parser.add_argument('--commit', help='Require this live commit before verification')
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--report-only', action='store_true', help='Record failures before a fix without failing the command')
    args = parser.parse_args()
    cases, counts = load_cases(args.source)
    client = app.test_client()
    def request(path, data=None, content_type=None):
        req = Request(args.base.rstrip('/') + path, data=data,
                      headers={'Content-Type': content_type} if content_type else {})
        with urlopen(req, timeout=60) as response:
            return response.read()
    release = json.loads(request('/version')) if args.base else client.get('/version').json
    if args.commit:
        assert release['commit'] == args.commit, release
    def check(case):
        payload = {'text': case['text']}
        result = json.loads(request('/api/parse', json.dumps(payload).encode(), 'application/json')) if args.base else client.post('/api/parse', json=payload).json
        return dict(id=case['id'], text=case['text'], errors=errors(case, result['items']))
    if args.base:
        with ThreadPoolExecutor(max_workers=2) as pool:
            checked = list(pool.map(check, cases))
    else:
        checked = [check(case) for case in cases]
    failures = [r for r in checked if r['errors']]
    outdir = args.output.parent
    outdir.mkdir(parents=True, exist_ok=True)
    # These IDs are sequential test IDs, unrelated to patient IDs in the source.
    text_buffer = io.StringIO(newline='')
    writer = csv.writer(text_buffer)
    writer.writerow(['patient_id', 'medication'])
    writer.writerows((c['id'], c['text']) for c in cases)
    raw = text_buffer.getvalue().encode('utf-8-sig')
    (outdir/'public_medication_input.csv').write_bytes(raw)
    if args.base:
        boundary = 'public-test-' + uuid.uuid4().hex
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="public_medication_input.csv"\r\nContent-Type: text/csv\r\n\r\n').encode() + raw + f'\r\n--{boundary}--\r\n'.encode()
        result_bytes = request('/upload', body, 'multipart/form-data; boundary=' + boundary)
    else:
        response = client.post('/upload', data={'file': (io.BytesIO(raw), 'public_medication_input.csv')})
        assert response.status_code == 200, response.data[:1000]
        result_bytes = response.data
    (outdir/'public_medication_result.xlsx').write_bytes(result_bytes)
    workbook = load_workbook(io.BytesIO(result_bytes))
    assert workbook.sheetnames == SHEETS, workbook.sheetnames
    audit_values = list(workbook['AuditTrail'].values)
    audit = [dict(zip(audit_values[0], row)) for row in audit_values[1:]]
    excel_errors = []
    for case in cases:
        rows = [r for r in audit if r['patient'] == case['id']]
        if len(rows) != 1:
            excel_errors.append(dict(id=case['id'], errors=[f'one medication split into {len(rows)} Excel rows']))
            continue
        row = rows[0]
        row['drug'] = row['parsed']
        row['conversions'] = []
        failed = errors(case, rows)
        if failed:
            excel_errors.append(dict(id=case['id'], errors=failed))
    report = dict(source_url=SOURCE_URL, source_sha256=SOURCE_SHA256, license='ODbL-1.0',
                  source_fields=FIELDS, patient_fields_uploaded=False, release=release,
                  **counts, api_passed=len(cases)-len(failures), api_failures=failures,
                  excel_passed=len(cases)-len(excel_errors), excel_failures=excel_errors,
                  sheets=workbook.sheetnames, clinical_validation=False)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps({k: report[k] for k in ('source_rows','selected_rows','unique_patterns','api_passed','excel_passed')}) + f' report={args.output}', flush=True)
    if not args.report_only:
        assert not failures and not excel_errors, 'See failure details in report'


if __name__ == '__main__':
    main()
