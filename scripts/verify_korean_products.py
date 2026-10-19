"""Verify exact Korean public product labels and synthetic dosing conditions."""
import argparse
from concurrent.futures import ThreadPoolExecutor
import csv
import hashlib
import io
import json
from pathlib import Path
import sys
from urllib.request import Request, urlopen
import uuid
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from app import app
from services.korean_products import CATALOG


def product_cases():
    for product in CATALOG['products']:
        variants=list(dict.fromkeys([product['name'],min(product['aliases'],key=len)]))
        for index,name in enumerate(variants):
            yield dict(product=product,id=product['code']+f'-{index}',
                       text=(name+' '+product['synthetic_frequency']).strip())


def errors_for(case, items):
    p=case['product']
    if len(items)!=1:
        return [f'expected one medication, received {len(items)}']
    item=items[0]
    errors=[]
    if item['drug']!=p['drug']:
        errors.append('ingredient mismatch')
    if item['route']!=('injection' if p['route']=='주사' else 'oral'):
        errors.append('route mismatch')
    if p['route']=='주사' and not p['synthetic_frequency']:
        if item['daily_dose_mg'] is not None or item['conversions'] or not item['needs_review']:
            errors.append('short-acting injection must remain unresolved')
    else:
        key='active_moiety_mg' if p['route']=='주사' else 'daily_dose_mg'
        if item.get(key)!=p['dose_mg']:
            errors.append('nominal active dose mismatch')
        if item['needs_review'] or not item['ok']:
            errors.append('valid synthetic dose unnecessarily held')
    return errors


def table(wb,name):
    values=list(wb[name].values)
    header=next(i for i,row in enumerate(values) if row[0] in ('patient_id','환자 ID','patient'))
    return [dict(zip(values[header],row)) for row in values[header+1:] if any(x is not None for x in row)]


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--base')
    parser.add_argument('--commit')
    parser.add_argument('--source',type=Path)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    def request(path, data=None, content_type=None):
        req=Request(args.base.rstrip('/')+path,data=data,
                    headers={'Content-Type':content_type} if content_type else {})
        with urlopen(req,timeout=120) as response:
            return response.read()
    if args.source:
        assert hashlib.sha256(args.source.read_bytes()).hexdigest()==CATALOG['sha256']
        source=load_workbook(args.source,read_only=True,data_only=True)
        original=list(source.active.values)
        for p in CATALOG['products']:
            row=original[p['source_row']-1]
            assert (str(row[8]),row[9],row[7],row[1])==(p['code'],p['name'],p['ingredient'],p['route'])
    client=app.test_client()
    version=json.loads(request('/version')) if args.base else client.get('/version').json
    if args.commit: assert version['commit']==args.commit,version
    cases=list(product_cases())
    def check(case):
        payload={'text':case['text'],'base_method':'CMD'}
        if args.base:
            data=json.loads(request('/api/parse',json.dumps(payload).encode(),'application/json'))
        else:
            data=client.post('/api/parse',json=payload).json
        return dict(id=case['id'],errors=errors_for(case,data['items']))
    if args.base:
        with ThreadPoolExecutor(max_workers=2) as pool: checked=list(pool.map(check,cases))
    else:
        checked=[check(case) for case in cases]
    args.output.parent.mkdir(parents=True,exist_ok=True)
    sheets=[]
    for dated in (False,True):
        buf=io.StringIO(newline=''); writer=csv.writer(buf)
        writer.writerow(['HID','PRESCR_DATE','DRUG','TABS_PER_DAY'] if dated else ['patient_id','medication'])
        writer.writerows([c['id'],'2026-10-19',c['text'],1] if dated else [c['id'],c['text']] for c in cases)
        raw=buf.getvalue().encode('utf-8-sig')
        name='korean_products_'+('dated' if dated else 'quick')
        (args.output.parent/(name+'.csv')).write_bytes(raw)
        local=client.post('/upload',data={'file':(io.BytesIO(raw),name+'.csv'),'base_method':'CMD'})
        assert local.status_code==200,local.data[:300]
        expected=load_workbook(io.BytesIO(local.data))
        content=local.data
        if args.base:
            boundary='ap-korean-'+uuid.uuid4().hex
            body=(f'--{boundary}\r\nContent-Disposition: form-data; name="base_method"\r\n\r\nCMD\r\n'
                  f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}.csv"\r\nContent-Type: text/csv\r\n\r\n').encode()+raw+f'\r\n--{boundary}--\r\n'.encode()
            content=request('/upload',body,'multipart/form-data; boundary='+boundary)
        (args.output.parent/(name+'.xlsx')).write_bytes(content)
        actual=load_workbook(io.BytesIO(content))
        assert actual.sheetnames==expected.sheetnames
        for sheet in actual.sheetnames:
            assert list(actual[sheet].values)==list(expected[sheet].values),sheet
        detail=table(actual,'보완근거')
        assert len(detail)==len(cases)
        by_id={row['환자 ID']:row for row in detail}
        assert set(by_id)=={case['id'] for case in cases}
        for case in cases:
            row=by_id[case['id']]
            assert row['환자 ID']==case['id'] and row['성분']==case['product']['drug']
            if case['product']['route']=='주사' and not case['product']['synthetic_frequency']:
                assert row['환산값 (mg/day)'] is None
            else:
                assert row['환산값 (mg/day)'] is not None,(case['id'],row)
        sheets.append(actual.sheetnames)
    failures=[r for r in checked if r['errors']]
    report=dict(products=len(CATALOG['products']),cases=len(cases),passed=len(cases)-len(failures),failures=failures,
                workbooks=sheets,release=version,source=CATALOG['source'],source_sha256=CATALOG['sha256'],
                conditions='Public product labels with synthetic QD/maintenance intervals; no patient data or clinical validation')
    args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:report[k] for k in ('products','cases','passed','failures')}))
    assert not failures


if __name__=='__main__': main()
