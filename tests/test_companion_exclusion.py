import io
import pandas as pd
import pytest
from openpyxl import load_workbook
from app import app
from services.exclusions import CATALOG, CLASSES, CLASS_LABELS
from services.frames import parse_frames, convert_frame, summarize_frames
from services.longitudinal_auto import automatic_analysis
from tests.workbook_helpers import records

COMPANIONS = [r for r in CATALOG['drugs'] if r['exclude']]


@pytest.mark.parametrize('row', COMPANIONS, ids=lambda r: r['ingredient'])
def test_source_ingredients_in_quick_and_dated_exports(row):
    name = row['ingredient']
    response = app.test_client().post('/api/export', json={'text': name + ' 50mg'})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert records(wb, 'Results')[0]['DDD (OLZ mg/day)'] == 0
    review = records(wb, 'Review')
    assert len(review) == 1
    assert review[0]['인식 성분'] == name
    assert review[0]['약물 분류'] == CLASS_LABELS[CLASSES[name]]
    assert review[0]['계산 처리'] == '배제 완료'
    assert review[0]['항정신병약 환산 기여값'] == 0
    audit = records(wb, 'AuditTrail')[0]
    assert audit['drug_class'] == CLASSES[name]
    assert row['source_cell'] in audit['exclusion_basis']
    assert audit['dose_mg'] == 50
    df = pd.DataFrame({'patient_id': ['001', '001', '001'],
                       'prescription_date': ['2020-04-15', '2020-04-16', '2020-04-16'],
                       'drug_name': [name, name + ' 50mg', 'risperidone 2mg tab'],
                       'tablets_per_day': ['', '', '1']})
    wb = load_workbook(automatic_analysis({'input': df}, ['DDD']))
    assert [r['DDD (OLZ mg/day)'] for r in records(wb, 'Results')] == [0, 4]
    review = records(wb, 'Review')
    assert len(review) == 2
    assert all(r['인식 성분'] == name and r['약물 분류'] == CLASS_LABELS[CLASSES[name]]
               and r['항정신병약 환산 기여값'] == 0 for r in review)
    audit = records(wb, 'AuditTrail')[0]
    assert audit['drug_class'] == CLASSES[name]
    assert row['source_cell'] in audit['exclusion_basis']


@pytest.mark.parametrize('text', ['lithiu 50mg', 'naltrexone (olanzapine) 5mg',
                                  'unknownxyz (modafinil) 50mg', 'clobazam + unknownxyz 5mg',
                                  'blonanserin 8mg', 'lithium 50mg, unknownxyz 5mg'])
def test_unknown_conflicts_and_missing_ap_factors_stay_reviewable(text):
    items = [convert_frame(f, ['DDD']) for f in parse_frames(text)]
    assert summarize_frames(items, ['DDD'])[0]['total_equivalent_dose_mg'] is None


def test_catalog_preserves_source_groups_and_antipsychotics():
    from services.converter import lookup
    assert len(CATALOG['drugs']) == 54 and len(COMPANIONS) == 39
    assert not set(CLASSES).intersection(lookup.source_drug)
    assert all(r['ingredient'] not in CLASSES for r in CATALOG['drugs'] if not r['exclude'])
    naltrexone = next(r for r in COMPANIONS if r['ingredient'] == 'naltrexone')
    assert naltrexone['source_group'] == 'stimulant'
    assert naltrexone['drug_class'] == 'opioid_antagonist'


@pytest.mark.parametrize('mode', ['csv', 'xlsx', 'structured'])
def test_uploads_exclude_all_source_companions_with_raw_fields(mode):
    names = [r['ingredient'] for r in COMPANIONS]
    frame = pd.DataFrame({'patient_id': [f'{i:03}' for i in range(len(names))],
                          'medication': [n + ' 50mg' for n in names]})
    if mode == 'structured':
        frame = pd.DataFrame({'patient_id': frame.patient_id, 'drug': names,
                              'dose': '', 'unit': '', 'frequency': 'unclear'})
    stream = io.BytesIO()
    if mode == 'xlsx':
        frame.to_excel(stream, index=False)
    else:
        stream.write(frame.to_csv(index=False).encode('utf-8'))
    stream.seek(0)
    response = app.test_client().post('/upload', data={'file': (stream, 'input.' + ('xlsx' if mode == 'xlsx' else 'csv'))})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert len(records(wb, 'Results')) == 39
    assert all(r['DDD (OLZ mg/day)'] == 0 for r in records(wb, 'Results'))
    review = records(wb, 'Review')
    assert len(review) == 39
    assert {r['인식 성분'] for r in review} == set(names)
    assert all(r['약물 분류'] and r['계산 처리'] == '배제 완료' and r['항정신병약 환산 기여값'] == 0 for r in review)
    assert len(records(wb, 'AuditTrail')) == 39
    assert all(r['source_text'] and r['exclusion_basis'] for r in records(wb, 'AuditTrail'))
