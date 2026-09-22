import io
import json

import pytest
from openpyxl import load_workbook
from app import app
from services import llm_candidates as llm, manual_suggestions as serving


def response(candidates=None, status='candidate_found'):
    return json.dumps(dict(status=status, candidates=candidates if candidates is not None else [
        dict(standard_name='aripiprazole', confidence='low', reason='Possible spelling variant')]))


@pytest.mark.parametrize('text', ['null', '[]', '{}', '{"candidates":null}',
    '{"status":"candidate_found","candidates":[null,"x",{}, {"standard_name":[]}]}',
    '```json\n{}\n```', '{} {}'])
def test_invalid_output_is_rejected(text):
    assert llm.validate_response(text, 'zzzzzz 2mg') == []


def test_closed_set_deduplication_limits_and_suffix():
    good = json.loads(response())['candidates'][0]
    candidates = [good, good, dict(good, standard_name='invented'),
                  dict(good, standard_name='risperidone'), dict(good, standard_name='olanzapine')]
    result = llm.validate_response(response(candidates), 'zzzzzz XR 2mg BID')
    assert len(result) == 2
    assert result[0]['replacement'] == 'aripiprazole XR 2mg BID'
    assert all(c['confirmed_drug'] is None and not c['auto_accepted'] for c in result)
    assert llm.validate_response(response(candidates, 'unknown'), 'zzzzzz 2mg') == []


@pytest.mark.parametrize('original', ['얀센 주사제 150mg', 'zzzzzz LAI 150mg',
    'zzzzzz 주사제 150mg', '환자 홍길동 zzzzzz 2mg', 'x' * 301])
def test_insufficient_or_unbounded_context_never_calls_model(monkeypatch, original):
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    monkeypatch.setattr(llm, '_ollama', lambda _: pytest.fail('Must not call backend'))
    assert llm.suggest_llm(original) == []


def test_disabled_and_existing_candidates_never_call_model(monkeypatch):
    monkeypatch.delenv('NAME_LLM_ENABLED', raising=False)
    monkeypatch.setattr(llm, '_ollama', lambda _: pytest.fail('Must not call backend'))
    assert llm.suggest_llm('zzzzzz 2mg') == []
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    assert serving.suggest_for_review('로핀 1mg')


def test_api_calls_model_once_without_converting_and_preserves_suffix(monkeypatch):
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    monkeypatch.setenv('NAME_LLM_BACKEND', 'ollama')
    calls = []
    def fake(messages):
        calls.append(messages)
        return response()
    monkeypatch.setattr(llm, '_ollama', fake)
    result = app.test_client().post('/api/parse', json={'text': 'zzzzzz 2mg BID'}).json
    assert len(calls) == 1
    assert '2mg' not in json.dumps(calls) and 'BID' not in json.dumps(calls)
    item = result['items'][0]
    assert item['drug'] is None and item['conversions'] == []
    candidate = item['suggestions'][0]
    confirmed = app.test_client().post('/api/parse', json={'text': candidate['replacement']}).json
    assert confirmed['items'][0]['drug'] == 'aripiprazole'
    assert confirmed['items'][0]['daily_dose_mg'] == 4


def test_backend_failure_is_private_and_nonfatal(monkeypatch, caplog):
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    def broken(_):
        raise TimeoutError('secret prescription and credential')
    monkeypatch.setattr(llm, '_ollama', broken)
    result = app.test_client().post('/api/parse', json={'text': 'zzzzzz 2mg'}).json
    assert result['items'][0]['suggestions'] == []
    assert 'secret' not in caplog.text and 'zzzzzz' not in caplog.text


def test_export_keeps_unconfirmed_candidates(monkeypatch):
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    monkeypatch.setattr(llm, '_ollama', lambda _: response())
    result = app.test_client().post('/api/export', json={'text': 'zzzzzz 2mg'})
    wb = load_workbook(io.BytesIO(result.data))
    audit = dict(zip(next(wb['AuditTrail'].values), list(wb['AuditTrail'].values)[1]))
    assert wb['Detailed'].max_row == 1
    assert json.loads(audit['name_candidates'])[0]['source'] == 'llm'


def test_windows_mlx_fails_without_import(monkeypatch):
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    monkeypatch.setenv('NAME_LLM_BACKEND', 'mlx')
    monkeypatch.setattr(llm.platform, 'system', lambda: 'Windows')
    assert llm.suggest_llm('zzzzzz 2mg') == []


@pytest.mark.parametrize('url', ['http://remote.example', 'https://remote.example',
    'http://user:password@localhost:11434', 'file:///tmp/model'])
def test_invalid_remote_configuration_is_rejected(monkeypatch, url):
    monkeypatch.setenv('NAME_LLM_URL', url)
    monkeypatch.delenv('NAME_LLM_API_KEY', raising=False)
    with pytest.raises(ValueError):
        llm._ollama([])


def test_ollama_request_contract_and_bounded_response(monkeypatch):
    monkeypatch.setenv('NAME_LLM_URL', 'http://127.0.0.1:11434')
    monkeypatch.setenv('NAME_LLM_MODEL', 'qwen2.5:7b')
    monkeypatch.delenv('NAME_LLM_API_KEY', raising=False)
    received = []
    class Reply(io.BytesIO):
        pass
    class Opener:
        def open(self, request, timeout):
            received.append((request, timeout))
            return Reply(json.dumps({'message': {'content': response()}}).encode())
    monkeypatch.setattr(llm, 'build_opener', lambda *_: Opener())
    assert llm._ollama([{'role': 'user', 'content': 'bounded name'}]) == response()
    request, timeout = received[0]
    assert request.full_url == 'http://127.0.0.1:11434/api/chat'
    payload = json.loads(request.data)
    assert payload['stream'] is False and payload['format'] == 'json'
    assert payload['model'] == 'qwen2.5:7b' and 1 <= timeout <= 30


def test_busy_backend_skips_without_queueing(monkeypatch):
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    monkeypatch.setattr(llm, '_ollama', lambda _: pytest.fail('Must not call backend'))
    with llm._lock:
        assert llm.suggest_llm('zzzzzz 2mg') == []
