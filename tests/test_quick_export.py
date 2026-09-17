import io
import pytest
from openpyxl import load_workbook
from app import app


@pytest.mark.parametrize('payload', [[], ['x'], {'text':None}, {'text':12}, {'text':''}, {'text':'a'*10001}])
@pytest.mark.parametrize('endpoint', ['/api/parse','/api/export'])
def test_bad_quick_input_returns_json_error(endpoint, payload):
    response=app.test_client().post(endpoint,json=payload)
    assert response.status_code == 400
    assert response.json['error']


def test_export_matches_parse_and_preserves_unresolved_rows():
    client=app.test_client()
    text='Sustenna 156mg monthly; risperidone 2mg QD; unknownxyz 5mg'
    parsed=client.post('/api/parse',json={'text':text}).json
    response=client.post('/api/export',json={'text':text})
    assert response.status_code == 200
    wb=load_workbook(io.BytesIO(response.data))
    totals=[dict(zip(next(wb['CellTotals'].values),r)) for r in list(wb['CellTotals'].values)[1:]]
    for actual, expected in zip(totals,parsed['totals']):
        assert actual['method']==expected['method']
        assert actual['total_equivalent_dose_mg']==expected['total_equivalent_dose_mg']
        assert actual['partial_equivalent_dose_mg']==expected['partial_equivalent_dose_mg']
    assert wb['AuditTrail'].max_row == 4
    assert wb['ReviewQueue'].max_row == 4
    assert 'VersionInfo' in wb.sheetnames


def test_quick_export_preserves_newlines_and_commas():
    text='risperidone 2mg QD,\nolanzapine 5mg HS'
    response=app.test_client().post('/api/export',json={'text':text})
    wb=load_workbook(io.BytesIO(response.data))
    assert wb['Results']['B2'].value == text
    assert wb['AuditTrail'].max_row == 3


@pytest.mark.parametrize('content', [
    'medication\n"Sustenna 156mg monthly; risperidone 2mg QD"\n',
    'patient_id;medication\nP1;"Sustenna 156mg monthly; risperidone 2mg QD"\n',
    'patient_id\tmedication\nP1\tSustenna 156mg monthly; risperidone 2mg QD\n',
])
def test_csv_delimiter_never_splits_inside_drug_name(content):
    response=app.test_client().post('/upload',data={'file':(io.BytesIO(content.encode()),'test.csv')})
    assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.data))
    rows=list(wb['AuditTrail'].values)
    assert len(rows)==3
    drug_column=rows[0].index('parsed')
    assert [r[drug_column] for r in rows[1:]]==['paliperidone','risperidone']


def test_empty_csv_returns_validation_error():
    response=app.test_client().post('/upload',data={'file':(io.BytesIO(b'\n'),'empty.csv')})
    assert response.status_code==400


def test_leading_equals_stays_text_in_export():
    response=app.test_client().post('/api/export',json={'text':'=risperidone 2mg QD'})
    assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.data))
    assert wb['Results']['B2'].value=="'=risperidone 2mg QD"
    assert wb['Results']['B2'].data_type=='s'
