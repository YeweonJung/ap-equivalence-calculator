import io
import pytest
from openpyxl import load_workbook
from app import app


@pytest.mark.parametrize('source', [
    'patient_id,medication\n001,risperidone 2mg QD\n002,unknownxyz 3mg\n',
    'HID,PRESCR_DATE,DRUG,TABS_PER_DAY\n001,2026-09-01,Risperidone 2mg tab,1\n002,2026-09-01,unknownxyz 3mg tab,1\n',
])
def test_all_sheets_are_plain_unfrozen_filterable_tables(source):
    response = app.test_client().post('/upload', data={'file': (io.BytesIO(source.encode()), 'synthetic.csv')})
    wb = load_workbook(io.BytesIO(response.data))
    for sheet in wb:
        assert sheet.freeze_panes is None
        assert sheet.sheet_view.pane is None
        assert not sheet.merged_cells.ranges
        assert sheet.auto_filter.ref.startswith('A1:')
        assert sheet.cell(1, 1).value
    assert wb['Results']['A2'].value == '001'
    headers = [c.value for c in wb['Review'][1]]
    assert '확인할 내용' in headers and '원문 약물' in headers
    assert '원본 행' in headers
    values = list(wb['Review'].values)
    assert any('unknownxyz' in str(row) for row in values[1:])
