import io
import json
import pytest
from openpyxl import load_workbook
from app import app
from services.name_distance import compare_letters
from services.drug_suggestions import suggest_drugs
from services.converter import convert_drug


@pytest.mark.parametrize('source,target,operation', [
    ('risperidnoe', 'risperidone', 'transpose'),
    ('risperidonee', 'risperidone', 'delete'),
    ('risperione', 'risperidone', 'insert'),
    ('rqsperidone', 'risperidone', 'replace'),
])
def test_character_alignment_and_confirmation(source, target, operation):
    difference = compare_letters(source, target)
    assert difference['distance'] == 1
    assert difference['edits'][0]['operation'] == operation
    response = app.test_client().post('/api/parse', json={'text': source+' 2mg BID'}).json
    item = response['items'][0]
    assert item['drug'] is None and not item['conversions']
    candidate = next(c for c in item['suggestions'] if c['alias'] == target)
    assert candidate['replacement'] == target+' 2mg BID'
    confirmed = app.test_client().post('/api/parse', json={'text': candidate['replacement']}).json['items'][0]
    assert confirmed['daily_dose_mg'] == 4


def test_different_ingredients_are_retained_without_automatic_selection():
    candidates = suggest_drugs('clopenthixl 5mg QD')
    assert any(c['drug'] == 'clopenthixol' for c in candidates)
    assert app.test_client().post('/api/parse', json={'text': 'clopenthixl 5mg QD'}).json['items'][0]['drug'] is None


def test_previously_registered_english_typo_also_requires_confirmation():
    item = app.test_client().post('/api/parse', json={'text':'risperidonne 2mg BID'}).json['items'][0]
    assert item['drug'] is None and not item['conversions']
    assert item['suggestions'][0]['alias'] == 'risperidone'


def test_formulation_schedule_and_unit_are_never_spell_corrected():
    candidate = next(c for c in suggest_drugs('quetiapnie XR 25mcg PRN') if c['alias'] == 'quetiapine')
    assert candidate['replacement'] == 'quetiapine XR 25mcg PRN'
    assert not suggest_drugs('qq 2mg')
    assert not suggest_drugs('xxxxxxxxxxxxx 5mg')


def test_export_keeps_unknown_original_character_evidence_and_no_total():
    response = app.test_client().post('/api/export', json={'text':'risperidnoe 2mg BID'})
    wb = load_workbook(io.BytesIO(response.data))
    audit = dict(zip(next(wb['AuditTrail'].values), list(wb['AuditTrail'].values)[1]))
    assert audit['original'] == 'risperidnoe 2mg BID' and audit['parsed'] is None
    assert json.loads(audit['name_candidates'])[0]['edits'][0]['operation'] == 'transpose'
    assert wb['Detailed'].max_row == 1
    assert 'name_candidates' in next(wb['ReviewQueue'].values)


@pytest.mark.parametrize('method,drug,dose', [
    ('CMD_DIRECT','risperidone',.27), ('CMD_DIRECT','quetiapine',27.64),
    ('CMD_INDIRECT','quetiapine',31.84), ('CMD_INDIRECT','chlorpromazine',28.77),
    ('CMD_INDIRECT','clozapine',39.96), ('CMD_INDIRECT','zotepine',16.35),
])
def test_published_table_1_olanzapine_equivalents(method, drug, dose):
    assert convert_drug(drug,dose,method) == pytest.approx(1)
    assert convert_drug('olanzapine',1,method,drug) == pytest.approx(dose)


def test_unpublished_cells_remain_unavailable_and_api_lists_nine_analyses():
    with pytest.raises(LookupError):
        convert_drug('chlorpromazine',100,'CMD_DIRECT')
    with pytest.raises(LookupError):
        convert_drug('paliperidone',9,'CMD_INDIRECT')
    result = app.test_client().post('/api/parse', json={'text':'risperidone 2.7mg QD'}).json
    assert len(result['totals']) == 9
    values = {c['method']:c['value'] for c in result['items'][0]['conversions']}
    assert values['CMD_DIRECT'] == values['CMD_INDIRECT'] == 10
