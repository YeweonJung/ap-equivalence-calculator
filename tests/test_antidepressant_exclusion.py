import io
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import app
from services.antidepressants import ALIASES, CATALOG, INGREDIENTS
from services.converter import available_methods
from services.frames import convert_frame, parse_frames, summarize_frames
from services.longitudinal import analyze, medication
from services.longitudinal_auto import automatic_analysis
from services.result_summary import value_column
from tests.test_longitudinal import prepared, rx
from tests.workbook_helpers import records


def converted(text):
    return [convert_frame(f, available_methods()) for f in parse_frames(text)]


@pytest.mark.parametrize('alias,ingredient', list(ALIASES.items()))
def test_registered_names_are_excluded_without_runtime_search(alias, ingredient):
    items = converted(alias + ' 50mg')
    assert len(items) == 1
    item = items[0]
    assert item['drug'] == ingredient
    assert item['drug_class'] == 'antidepressant'
    assert item['status'] == 'non_target' and not item['needs_review']
    assert item['dose_mg'] == 50  # The source dose is not overwritten with zero.
    assert len(item['conversions']) == len(available_methods())
    assert all(c['value'] == 0 for c in item['conversions'])
    totals = summarize_frames(items, available_methods())
    assert all(t['total_equivalent_dose_mg'] == 0 and not t['needs_review'] for t in totals)
    canonical, _, _, issues, kind = medication(alias + ' 50mg tab', '')
    assert canonical == ingredient and kind == 'non_target' and issues == []


def test_catalog_sources_and_no_ap_ingredients():
    from services.converter import lookup
    assert len(INGREDIENTS) == 18
    assert not INGREDIENTS.intersection(lookup.source_drug)
    assert all(row['source_ids'] and all(s in CATALOG['sources'] for s in row['source_ids'])
               for row in CATALOG['drugs'])


@pytest.mark.parametrize('text', ['sertraline', '렉사프로 10', 'Wellbutrin 150mg PRN'])
def test_exclusion_does_not_require_a_dose_unit_or_schedule(text):
    items = converted(text)
    assert len(items) == 1 and items[0]['status'] == 'non_target'
    assert not items[0]['needs_review']
    assert summarize_frames(items, ['DDD'])[0]['total_equivalent_dose_mg'] == 0


@pytest.mark.parametrize('text', [
    'sertralin 50mg', 'Symbyax 6/25mg', 'unknownxyz 5mg',
    'fluoxetine olanzapine 5mg', 'fluoxetine (olanzapine) 5mg',
    'unknownxyz (sertraline) 50mg', 'sertraline unknownxyz 50mg',
    'fluoxetine / unknownxyz', 'sertraline + unknownxyz 5mg',
    'sertraline 50mg, unknownxyz 5mg', 'sertraline 50mg, blonanserin 8mg',
])
def test_unknown_conflicts_and_missing_ap_factors_never_become_zero(text):
    assert summarize_frames(converted(text), ['DDD'])[0]['total_equivalent_dose_mg'] is None


def test_mixed_regimen_and_other_non_targets_unchanged():
    mixed = converted('sertraline 50mg, risperidone 2mg QD, lithium 600mg')
    assert summarize_frames(mixed, ['DDD'])[0]['total_equivalent_dose_mg'] == 120
    assert mixed[-1]['conversions'] == []
    assert summarize_frames(converted('lithium 600mg'), ['DDD'])[0]['total_equivalent_dose_mg'] is None


@pytest.mark.parametrize('drug,product', [
    ('fluoxetine olanzapine 5mg tab', ''),
    ('sertraline 50mg tab', 'Risperdal 2mg'),
    ('sertraline (olanzapine) 5mg tab', ''),
    ('unknownxyz 5mg tab', 'Zoloft 50mg'),
])
def test_dated_conflicts_do_not_silently_exclude(drug, product):
    rows = prepared([rx(drug=drug, product=product)], 'prescription_date')
    result = analyze(rows, [('P001', date(2020, 4, 15))], ['DDD'], 'prescription_date')[0][0]
    assert result['equivalent_mg'] is None and result['status'] == 'review'


def test_exact_date_totals_audit_and_no_carryover():
    source = pd.DataFrame([
        rx(day='2020-04-15', patient='001', drug='졸로푸트정 50mg', tabs=''),
        rx(day='2020-04-16', patient='001', drug='Risperidone 2mg tab', tabs='1'),
        rx(day='2020-04-16', patient='001', drug='심발타 30mg', tabs='bad'),
        rx(day='2020-04-17', patient='001', drug='unknownxyz 5mg tab', tabs='1'),
        rx(day='2020-04-17', patient='001', drug='렉사프로 10mg', tabs='1'),
    ])
    wb = load_workbook(automatic_analysis({'input': source}, ['DDD']))
    assert wb.sheetnames == ['Results', 'MedicationResults', 'Review', 'AuditTrail']
    totals = records(wb, 'Results')
    assert [r['DDD (CPZ mg/day)'] for r in totals] == [0, 120, None]
    assert [r['patient_id'] for r in totals] == ['001'] * 3
    details = records(wb, 'MedicationResults')
    assert len(details) == 5
    excluded = [r for r in details if r['결과 상태'] == '항우울제 제외 (0)']
    assert len(excluded) == 3
    assert all(r['DDD (CPZ mg/day)'] == 0 for r in excluded)
    assert len(records(wb, 'Review')) == 1  # Unknown drug only.
    audit = records(wb, 'AuditTrail')
    assert len(audit) == 5
    assert audit[0]['drug_class'] == 'antidepressant'
    assert audit[0]['exclusion_basis'] and audit[0]['daily_mg'] is None
    assert audit[0]['drug'] == source.iloc[0]['DRUG']


@pytest.mark.parametrize('mode', ['csv', 'xlsx', 'quick', 'structured'])
def test_all_static_export_paths_keep_zero_and_audit_without_review(mode):
    client = app.test_client()
    if mode == 'quick':
        response = client.post('/api/export', json={'text': 'Zoloft 50mg'})
    else:
        frame = pd.DataFrame({'patient_id': ['001'], 'medication': ['Zoloft 50mg']})
        if mode == 'structured':
            frame = pd.DataFrame({'patient_id': ['001'], 'drug': ['Zoloft'],
                                  'dose': [''], 'unit': [''], 'frequency': ['unclear']})
        stream = io.BytesIO()
        if mode == 'xlsx':
            frame.to_excel(stream, index=False)
        else:
            stream.write(frame.to_csv(index=False).encode('utf-8'))
        stream.seek(0)
        response = client.post('/upload', data={'file': (stream, 'input.' + ('xlsx' if mode == 'xlsx' else 'csv'))})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames == ['Results', 'MedicationResults', 'Review', 'AuditTrail']
    assert records(wb, 'Results')[0]['DDD (CPZ mg/day)'] == 0
    assert records(wb, 'MedicationResults')[0][value_column('DDD')] == 0
    assert not records(wb, 'Review')
    audit = records(wb, 'AuditTrail')[0]
    assert audit['drug_class'] == 'antidepressant' and audit['status'] == 'non_target'
    assert audit['exclusion_basis'] and audit['source_text']
