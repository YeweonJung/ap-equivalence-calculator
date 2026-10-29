"""Regressions discovered using medication-only MIMIC-IV demo patterns.

https://physionet.org/content/mimic-iv-demo/2.2/ (ODbL 1.0).
No patient identifiers, dates, or other clinical fields are included.
"""
import io
import pytest
from openpyxl import load_workbook
from app import app
from tests.workbook_helpers import records


def parse(text):
    response = app.test_client().post('/api/parse', json={'text': text})
    assert response.status_code == 200
    return response.json['items']


@pytest.mark.parametrize('text,expected', [
    ('QUEtiapine Fumarate 100mg PO/NG 2 times a day', 200),
    ('Haloperidol 1mg PO 2 times a day', 2),
    ('OLANZapine (Disintegrating Tablet) 5mg PO 2 times a day', 10),
    ('ZIPRASidone Hydrochloride 40mg PO/NG 2 times a day', 80),
    ('Perphenazine 8mg PO/NG 3 times a day', 24),
])
def test_route_and_frequency_remain_in_same_medication(text, expected):
    items = parse(text)
    assert len(items) == 1
    item = items[0]
    assert item['daily_dose_mg'] == expected
    assert item['needs_review'] is False
    assert item['original'] == text[item['source_start']:item['source_end']] == text


@pytest.mark.parametrize('route', ['IV', 'IM', 'SC', 'intramuscular', 'subcutaneous', 'injectable'])
def test_injection_before_numeric_frequency_never_becomes_oral(route):
    items = parse(f'Haloperidol 2.5mg {route} 6 times a day')
    assert len(items) == 1
    assert items[0]['route'] == 'injection'
    assert items[0]['daily_dose_mg'] is None
    assert not items[0]['conversions']
    assert items[0]['needs_review']


def test_zero_frequency_is_not_lost_and_missing_frequency_stays_reviewable():
    invalid = parse('QUEtiapine Fumarate 25mg PO/NG 0 times a day')
    assert len(invalid) == 1
    assert invalid[0]['daily_dose_mg'] is None and not invalid[0]['conversions']
    missing = parse('QUEtiapine Fumarate 25mg PO/NG')[0]
    assert missing['frequency'] == 'ASSUMED_QD' and missing['needs_review']


def test_route_fix_keeps_unknown_drug_boundaries():
    text = 'quetiapine 100mg PO/NG 2 times a day Unknownzapine 7mg olanzapine 5mg PO 1 times a day'
    items = parse(text)
    assert [item['drug'] for item in items] == ['quetiapine', None, 'olanzapine']
    assert [item['daily_dose_mg'] for item in items] == [200, None, 5]
    for item in items:
        assert text[item['source_start']:item['source_end']] == item['original']


def test_public_pattern_upload_and_quick_download_match_safe_api_values():
    client = app.test_client()
    source = b'patient_id,medication\nPUBLIC001,QUEtiapine Fumarate 100mg PO/NG 2 times a day\nPUBLIC002,Haloperidol 2.5mg IV 6 times a day\n'
    response = client.post('/upload', data={'file': (io.BytesIO(source), 'public-patterns.csv')})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    rows = records(wb, 'MedicationResults')
    assert len(rows) == 2
    assert rows[0]['daily_dose_mg'] == 200
    assert rows[1]['daily_dose_mg'] is None
    assert rows[1]['status'] == 'unsupported_formulation'
    response = client.post('/api/export', json={'text': 'QUEtiapine Fumarate 100mg PO/NG 2 times a day'})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert len(wb.sheetnames) == 7
    assert records(wb, 'MedicationResults')[0]['daily_dose_mg'] == 200
