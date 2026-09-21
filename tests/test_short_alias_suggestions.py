import io
import json

import pytest
from openpyxl import load_workbook
from app import app
from services.manual_suggestions import suggest_for_review
from services.short_alias_suggestions import short_alias_candidates


def test_reported_input_requires_selection_and_preserves_dose_frequency():
    client = app.test_client()
    item = client.post('/api/parse', json={'text': '할리l 2mg BID'}).json['items'][0]
    assert item['drug'] is None and not item['conversions']
    candidate = next(c for c in item['suggestions'] if c['alias'] == '할돌')
    assert candidate['replacement'] == '할돌 2mg BID'
    assert candidate['confidence'] == 'low' and not candidate['policy_eligible']
    assert not candidate['auto_accepted'] and candidate['confirmed_drug'] is None
    assert '낮은 유사도' in candidate['explanation']
    fixed = client.post('/api/parse', json={'text': candidate['replacement']}).json['items'][0]
    assert fixed['drug'] == 'haloperidol' and fixed['daily_dose_mg'] == 4
    assert fixed['conversions']


@pytest.mark.parametrize('text', ['할', '할리ll', '할리abc', '가리l', '할리 l',
    'abc', '할리l;리튬', '할리l\n리튬', '123', '할리<', '할리;', '리스'])
def test_fallback_rejects_inputs_outside_narrow_gate(text):
    assert short_alias_candidates(text) == []


def test_two_syllable_aliases_are_not_special_cased_to_haldol():
    assert short_alias_candidates('리틈 100mg')[0]['alias'] == '리튬'


def test_rollback_disables_new_fallback(monkeypatch):
    monkeypatch.setenv('NAME_RETRIEVAL_ENABLED', '0')
    assert suggest_for_review('할리l 2mg BID') == []


def test_excel_records_unconfirmed_short_alias_candidate():
    response = app.test_client().post('/api/export', json={'text': '할리l 2mg BID'})
    wb = load_workbook(io.BytesIO(response.data))
    rows = list(wb['AuditTrail'].values)
    audit = dict(zip(rows[0], rows[1]))
    assert audit['parsed'] is None and wb['Detailed'].max_row == 1
    candidate = json.loads(audit['name_candidates'])[0]
    assert candidate['alias'] == '할돌' and not candidate['auto_accepted']
