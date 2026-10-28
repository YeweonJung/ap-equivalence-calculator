import io
import math

import pytest
from openpyxl import load_workbook

from app import app
from services.converter import convert_drug, normalize_target
from services.frames import convert_frame, parse_frames
from services.result_summary import METHOD_ORDER
from services.supplemented import quick_cases


def calculate(text, methods=METHOD_ORDER):
    items = [convert_frame(f, methods) for f in parse_frames(text)]
    return quick_cases(items, methods)


def test_base_values_fixed_and_first_available_supplement_only():
    cases = calculate('olanzapine 5mg QD, paliperidone 6mg QD', ['CMD'])
    assert [c['case_id'] for c in cases] == ['CMD+DDD']
    for case in cases:
        base, alternative = case['details']
        assert base['actual'] == 'CMD' and base['value'] == 5
        assert alternative['actual'] == case['complement']
        assert case['methods'] == 'CMD + ' + case['complement']
        assert case['total'] == pytest.approx(round(5 + convert_drug('paliperidone', 6, case['complement'], 'olanzapine'), 4))
        assert case['missing'] == 0 and case['supplemented'] == 1
        assert all(r['case_id'] == case['case_id'] for r in case['details'])


def test_missing_drugs_are_filled_in_order_without_replacing_earlier_values():
    cases = calculate('olanzapine 5mg QD, paliperidone 6mg QD, clopenthixol 60mg QD', ['CMD'])
    result, = cases
    assert result['case_id'] == 'CMD+DDD+GARDNER'
    assert [r['actual'] for r in result['details']] == ['CMD', 'DDD', 'GARDNER']
    assert result['details'][0]['value'] == 5
    assert result['details'][1]['value'] == 10
    expected = 15 + convert_drug('clopenthixol', 60, 'GARDNER', 'olanzapine')
    assert result['total'] == pytest.approx(round(expected, 4))
    assert result['partial'] is None and result['missing'] == 0


def test_complete_base_is_not_repeated_for_every_fallback():
    cases = calculate('olanzapine 5mg QD, risperidone 2mg QD', ['CMD', 'MED'])
    assert [c['case_id'] for c in cases] == ['CMD', 'MED']
    assert all(c['complement'] == '' and c['supplemented'] == 0 for c in cases)


@pytest.mark.parametrize('text', ['unknownxyz 3mg', 'aripiprazole LAI 400mg', 'paliperidone 6'])
def test_input_and_formulation_blockers_never_create_complete_alternatives(text):
    cases = calculate(text, ['CMD'])
    assert len(cases) == 1 and cases[0]['case_id'] == 'CMD'
    assert cases[0]['total'] is None


@pytest.mark.parametrize('text', ['aripiprazole LAI 400mg q4w', 'risperidone Okedi 100mg q4w'])
def test_no_oral_bridge_permits_only_an_explicit_validated_depot_ddd_case(text):
    from services.lai_support import convert_injection
    item = convert_frame(parse_frames(text)[0], ['CMD', 'DDD'])
    assert item['oral_equivalent_mg'] is None
    assert next(c for c in item['conversions'] if c['method'] == 'CMD')['value'] is None
    cases = quick_cases([item], ['CMD'])
    assert len(cases) == 1 and cases[0]['case_id'] == 'CMD+DDD'
    assert cases[0]['total'] == pytest.approx(round(convert_injection(item, 'DDD', 'olanzapine'), 4))
    detail, = cases[0]['details']
    assert detail['actual'] == 'DDD' and '제형별 DDD' in detail['basis']
    assert '제형별 DDD로 보완' in detail['reason']
    assert detail['factor'] * detail['dose'] == pytest.approx(detail['value'])


def test_no_supported_fallback_and_excluded_only_have_one_case_per_base():
    unknown = calculate('blonanserin 8mg QD', ['CMD'])
    assert len(unknown) == 1 and unknown[0]['total'] is None
    excluded = calculate('lithium 600mg', ['CMD'])
    assert len(excluded) == 1 and excluded[0]['total'] == 0


def test_same_target_is_used_for_all_drugs_in_every_case():
    cases = calculate('paliperidone 6mg QD, clopenthixol 60mg QD')
    for case in cases:
        target = normalize_target(case['base'])
        assert {r['target'] for r in case['details']} == {case['target']}
        values = []
        for detail in case['details']:
            if detail['value'] is not None:
                expected = convert_drug(detail['drug'], detail['dose'], detail['actual'], target)
                assert detail['value'] == pytest.approx(expected)
                values.append(expected)
        if case['total'] is not None:
            assert case['total'] == pytest.approx(round(math.fsum(values), 4))
        else:
            assert case['missing'] > 0


def test_quick_api_and_excel_cases_and_medication_methods_match_exactly():
    text = 'olanzapine 5mg QD, paliperidone 6mg QD, unknownxyz 2mg'
    client = app.test_client()
    data = client.post('/api/parse', json={'text': text}).json
    assert data['case_policy'] == 'selected-base-sequential-3.1'
    assert data['base_method'] == 'CMD' and len(data['cases']) == 1
    response = client.post('/api/export', json={'text': text})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames[:2] == ['보완합계', '보완근거']
    def rows(name):
        values = list(wb[name].values)
        return [dict(zip(values[0], r)) for r in values[1:]]
    summary = {r['계산안']: r for r in rows('보완합계')}
    assert len(summary) == len(data['cases'])
    assert '성분' not in summary[next(iter(summary))]
    detail_rows = rows('보완근거')
    for case in data['cases']:
        row = summary[case['case_id']]
        assert row['보완 총합 (mg/day)'] == case['total']
        assert row['부분합 (mg/day)'] == case['partial']
        assert (row['보완 방법'] or '') == case['complement']
        assert row['미환산 수'] == case['missing']
        assert [(r['성분'], r['적용 방법'] or '') for r in detail_rows] == [
            (r['drug'], r['actual']) for r in case['details']]


def test_date_and_sheet_provenance_with_repeated_source_row_numbers():
    import pandas as pd
    from services.longitudinal_auto import automatic_analysis
    from tests.test_longitudinal import rx
    sheets = {'First': pd.DataFrame([rx(patient='001', drug='Paliperidone 6mg tab', tabs='1')]),
              'Second': pd.DataFrame([rx(patient='001', drug='Olanzapine 5mg tab', tabs='1'),
                                      rx(patient='002', day='2020-05-01', drug='Olanzapine 5mg tab', tabs='1')])}
    wb = load_workbook(automatic_analysis(sheets, ['CMD']))
    data = list(wb['보완합계'].values)
    summaries = [dict(zip(data[0], row)) for row in data[1:]]
    assert len([r for r in summaries if r['환자 ID'] == '001']) == 1
    assert next(r for r in summaries if r['환자 ID'] == '001' and r['계산안'] == 'CMD+DDD')['보완 총합 (mg/day)'] == 15
    data = list(wb['보완근거'].values)
    detail = [dict(zip(data[0], row)) for row in data[1:]]
    cmd_med = [r for r in detail if r['환자 ID'] == '001']
    assert {(r['원본 시트'], r['원본 행']) for r in cmd_med} == {('First', 2), ('Second', 2)}
