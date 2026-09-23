import io
import json
import zipfile
from datetime import date

import pandas as pd
import pytest

from services.longitudinal import (analyze, detect_mapping, prepare, reference_pairs,
                                  safe_csv, parse_date, medication, is_longitudinal)


def rx(day='2020-04-15', days='14', tabs='0.5', drug='Quetiapine 25mg tab', product='', patient='P001'):
    return dict(HID=patient, PRESCR_DATE=day, PRESCR_DAYS=days, TABS_PER_DAY=tabs, DRUG=drug, COMMERCIAL_DRUG=product)


def prepared(rows, policy='review'):
    frame = pd.DataFrame(rows).fillna('')
    return prepare(frame, detect_mapping(frame.columns), policy)


def result(records, day, policy='review'):
    return analyze(records, [('P001', date.fromisoformat(day))], ['DDD'], policy)[0][0]


def test_127_128_synthetic_regimen_replacement_and_no_resurrection():
    rows = [rx(), rx('2020-04-26', '2', '1')]
    records = prepared(rows, 'replace')
    assert records[0]['effective_end'] == date(2020, 4, 26)
    assert result(records, '2020-04-25', 'replace')['status'] == 'calculated_assumption'
    before = result(records, '2020-04-25')['equivalent_mg']
    assert result(records, '2020-04-26')['equivalent_mg'] == before * 2
    assert result(records, '2020-04-28')['status'] == 'no_record'


def test_default_blocks_only_overlap_window():
    records = prepared([rx(), rx('2020-04-26', '14', '1')])
    assert result(records, '2020-04-25')['status'] == 'calculated'
    assert result(records, '2020-04-26')['equivalent_mg'] is None
    assert result(records, '2020-04-29')['status'] == 'calculated'


def test_end_inclusive_no_gap_fill():
    records = prepared([rx(), rx('2020-05-01', '3', '1')], 'replace')
    assert result(records, '2020-04-28')['status'] == 'calculated'
    assert result(records, '2020-04-29')['status'] == 'no_record'
    assert result(records, '2020-05-04')['status'] == 'no_record'


@pytest.mark.parametrize('rows', [[rx(), rx()], [rx(), rx(tabs='1')], [rx(), rx(drug='Quetiapine 100mg tab')]])
def test_same_day_multiple_not_summed(rows):
    records = prepared(rows, 'replace')
    assert all('same_day_multiple_orders' in r['issues'] for r in records)
    assert result(records, '2020-04-15')['equivalent_mg'] is None


def test_different_ingredients_sum():
    records = prepared([rx(), rx(drug='Risperidone 2mg tab', tabs='1')])
    assert result(records, '2020-04-15')['status'] == 'calculated'


def test_formulation_not_replaced():
    records = prepared([rx(), rx('2020-04-20', drug='Quetiapine 25mg ER tab')], 'replace')
    assert records[0]['effective_end'] == date(2020, 4, 29)
    assert result(records, '2020-04-20')['status'] == 'review'


@pytest.mark.parametrize('days', ['0', '-1', '1.5', '', 'inf', 'nan', '4000'])
def test_invalid_days_block_not_zero(days):
    records = prepared([rx(days=days)])
    assert 'invalid_days' in records[0]['issues']
    assert result(records, '2020-04-15')['equivalent_mg'] is None


@pytest.mark.parametrize('tabs', ['0', '-1', '', 'nan', 'inf', '900'])
def test_bad_daily_dose(tabs):
    assert 'invalid_daily_tablets' in prepared([rx(tabs=tabs)])[0]['issues']


def test_invalid_new_order_blocks_old_dose():
    records = prepared([rx(), rx('2020-04-20', days='0')], 'replace')
    assert records[0]['effective_end'] == date(2020, 4, 20)
    assert result(records, '2020-04-20')['status'] == 'review'


def test_mismatched_lai_never_uses_oral_20():
    r = prepared([rx(drug='Aripiprazole 20mg', product='아빌리파이메인테나주 300mg/1.5ml(Aripiprazole)')])[0]
    assert 'strength_conflict' in r['issues']
    assert 'injection_requires_review' in r['issues']
    assert result([r], '2020-04-15')['equivalent_mg'] is None


def test_matching_invega():
    r = prepared([rx(drug='Paliperidone 3mg ER tab', product='인베가서방정 3mg(Paliperidone)', tabs='2')])[0]
    assert r['issues'] == []
    assert r['daily_mg'] == 6


def test_non_target_and_unknown_not_conflated():
    assert result(prepared([rx(drug='Escitalopram 10mg tab')]), '2020-04-15')['status'] == 'no_record'
    assert result(prepared([rx(drug='unknown 10mg tab')]), '2020-04-15')['status'] == 'review'


def test_invalid_date_blocks_all_references():
    assert result(prepared([rx(day='uncertain')]), '2020-01-01')['status'] == 'review'


def test_column_aliases_and_ambiguity():
    assert is_longitudinal(['처방일', '처방일수'])
    assert is_longitudinal(['PRESCR_DATE', '처방일', 'DRUG'])
    assert detect_mapping(['HID', 'patient_id'])['patient'] == ''
    assert detect_mapping(['patient id', 'prescription_date', 'drug_name', 'daily_tablets', 'days_supply'])['daily'] == 'daily_tablets'


def test_strict_dates():
    assert parse_date('20200229') == parse_date('2020/2/29')
    with pytest.raises(ValueError):
        parse_date('02/03/2020')


def test_reference_pairs_missing_patient_and_duplicates():
    records = prepared([rx(patient='001'), rx(patient='002')])
    refs = pd.DataFrame([dict(patient_id='003', reference_date='2020-04-15')])
    pairs = reference_pairs(records, 'per_patient', reference_frame=refs)
    assert pairs[0][0] == '003'
    assert analyze(records, pairs, ['DDD'])[0][0]['status'] == 'no_record'
    with pytest.raises(ValueError):
        reference_pairs(records, 'per_patient', reference_frame=pd.concat([refs, refs]))


def test_formula_injection_escaped():
    assert "'=HYPERLINK" in safe_csv([dict(patient_id='=HYPERLINK(1)')]).decode('utf-8-sig')


def test_api_inspection_download_and_upload_guard():
    from app import app
    client = app.test_client()
    raw = pd.DataFrame([rx(patient='001')]).to_csv(index=False).encode('utf-8-sig')
    response = client.post('/api/longitudinal/inspect', data={'file':(io.BytesIO(raw), 'rx.csv')})
    assert response.json['recognized'] is True
    assert '001' not in response.get_data(as_text=True)
    response = client.post('/api/longitudinal/analyze', data={'file':(io.BytesIO(raw), 'rx.csv'), 'ack':'yes', 'reference_date':'2020-04-15'})
    assert response.status_code == 200
    assert response.headers['Cache-Control'] == 'no-store'
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(response.data), read_only=True)
    assert wb.sheetnames == ['Results', 'MedicationResults', 'Review']
    assert list(wb['Results'].values)[1][0] == '001'
    assert json.loads(wb.properties.description)['rule_version'] == 'prescription-date-2.0'
    wb.close()
    response = client.post('/upload', data={'file':(io.BytesIO(raw), 'rx.csv')})
    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    auto = load_workbook(io.BytesIO(response.data), read_only=True)
    assert json.loads(auto.properties.description)['mode'] == 'automatic'
    auto.close()


def test_manual_mapping_and_cp949():
    from app import app
    client = app.test_client()
    frame = pd.DataFrame([rx()]).rename(columns={'HID':'연구코드'})
    mapping = detect_mapping(frame.columns)
    mapping['patient'] = '연구코드'
    response = client.post('/api/longitudinal/analyze', data={'file':(io.BytesIO(frame.to_csv(index=False).encode('cp949')), 'rx.csv'), 'mapping':json.dumps(mapping), 'ack':'yes', 'mode':'all_dates'})
    assert response.status_code == 200


def test_blank_records_preserve_row_numbers():
    from services.longitudinal_api import read_csv_upload
    from werkzeug.datastructures import FileStorage
    text = pd.DataFrame([rx(), rx('2020-04-26')]).to_csv(index=False)
    text = text.replace('\n', '\n\n', 1)
    frame, _ = read_csv_upload(FileStorage(io.BytesIO(text.encode()), filename='rx.csv'))
    assert [r['source_row'] for r in prepare(frame, detect_mapping(frame.columns))] == [3, 4]
