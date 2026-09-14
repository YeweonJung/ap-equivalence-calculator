import io
import pytest
from openpyxl import load_workbook
from app import app
from services.converter import convert_drug


def parse(text):
    response = app.test_client().post('/api/parse', json={'text': text})
    assert response.status_code == 200
    return response.get_json()


@pytest.mark.parametrize('text,drug,oral', [
    ('paliperidone 25mg PP1M', 'paliperidone', 3),
    ('paliperidone 75mg PP1M', 'paliperidone', 6),
    ('paliperidone 100mg PP1M', 'paliperidone', 9),
    ('paliperidone 150mg PP1M', 'paliperidone', 12),
    ('Trevicta 263mg', 'paliperidone', 6),
    ('paliperidone 350mg PP3M', 'paliperidone', 9),
    ('paliperidone 700mg PP6M', 'paliperidone', 9),
    ('Byannli 1000mg', 'paliperidone', 12),
    ('olanzapine 150mg LAI q2w', 'olanzapine', 10),
    ('Zypadhera 405mg q4w', 'olanzapine', 15),
    ('olanzapine 300mg LAI q2w', 'olanzapine', 20),
    ('Okedi 75mg', 'risperidone', 3),
])
def test_product_label_oral_bridge_is_separate_from_ddd(text, drug, oral):
    data = parse(text)
    assert len(data['items']) == 1
    item = data['items'][0]
    assert item['drug'] == drug and item['oral_equivalent_mg'] == oral
    assert item['daily_dose_mg'] != oral  # mass/day is not the oral bridge
    for c in item['conversions']:
        if c['method'] == 'DDD':
            assert c['basis'] == 'WHO depot DDD'
        else:
            try:
                expected = round(convert_drug(drug, oral, c['method']), 4)
            except LookupError:
                expected = None
            assert c['value'] == expected
            if expected is not None:
                assert 'LAI 직접 환산 아님' in c['basis']
        total = next(t for t in data['totals'] if t['method'] == c['method'])
        assert total['total_equivalent_dose_mg'] == c['value']


@pytest.mark.parametrize('text', [
    'aripiprazole 400mg LAI 1.5개월', 'aripiprazole 400mg (LAI 1.5개월)',
    'aripiprazole 400mg LAI q0.5w', 'aripiprazole 400mg LAI q-4w',
    'aripiprazole 400mg LAI monthly BID', 'aripiprazole 400mg PP1M',
    'paliperidone 156mg PP1M', 'paliperidone palmitate 156mg PP1M',
    'paliperidone 100mg PP3M', 'paliperidone 100mg PP6M',
    'paliperidone 100mg PP1M PP3M', 'olanzapine 300mg LAI q3w',
    'aripiprazole lauroxil 441mg LAI q4w', 'Maintena 960mg LAI q8w',
    'risperidone 75mg LAI q4w', 'Okedi 75mg BID',
    'Zypadhera 405mg q4w loading',
])
def test_ambiguous_lai_has_no_values_or_wrong_oral_partial(text):
    data = parse(text)
    assert all(c['value'] is None for i in data['items'] for c in i['conversions'])
    assert all(t['partial_equivalent_dose_mg'] is None and t['total_equivalent_dose_mg'] is None for t in data['totals'])


def test_lai_annotation_stays_attached_and_missing_unit_is_auditable():
    for text in ('aripiprazole 400mg LAI 1개월', 'aripiprazole 400 LAI 1개월'):
        item = parse(text)['items'][0]
        assert item['route'] == 'injection' and item['interval_days'] == 30
        assert item['daily_dose_mg'] == pytest.approx(400/30)
        assert len(parse(text)['items']) == 1
    assert parse('aripiprazole 400 LAI 1개월')['items'][0]['unit_assumed']


def test_two_month_aripiprazole_uses_product_56_days():
    a = parse('aripiprazole 960mg LAI every 2 months')['items'][0]
    b = parse('aripiprazole 960mg LAI q8w')['items'][0]
    assert a['interval_days'] == b['interval_days'] == 56
    assert a['conversions'] == b['conversions']
    assert a['oral_equivalent_mg'] is None


def test_mixed_oral_lai_totals_use_each_methods_own_basis():
    data = parse('paliperidone 100mg PP1M, olz 5mg')
    totals = {t['method']: t for t in data['totals']}
    assert totals['MED']['total_equivalent_dose_mg'] == 27.5
    assert totals['DDD']['total_equivalent_dose_mg'] == 550
    assert totals['CMD']['total_equivalent_dose_mg'] is None
    assert totals['CMD']['unresolved_count'] == 1


def test_export_labels_estimates_and_preserves_blank_unsupported_methods():
    csv = 'patient_id,medication\nP1,"paliperidone 100mg PP1M, olz 5mg"\nP2,aripiprazole 400mg LAI 1.5개월\n'
    response = app.test_client().post('/upload', data={'method': 'ALL', 'file': (io.BytesIO(csv.encode()), 'synthetic.csv')})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    result = list(wb['Results'].values)
    assert result[1][4] == 22.5 and result[1][9] == 27.5
    assert result[1][3] is None and result[1][8] is None
    assert '경구 대응용량 기반 추정' in result[1][13]
    assert all(v is None for v in result[3][3:13])
    detail = [dict(zip(next(wb['Detailed'].values), r)) for r in list(wb['Detailed'].values)[1:]]
    med = next(r for r in detail if r['drug'] == 'paliperidone' and r['method'] == 'MED')
    assert med['oral_equivalent_mg'] == 9 and '경구' in med['conversion_basis']
    assert '7652' in med['conversion_source']
    assert len(list(wb['InjectionInfo'].values)) > 10


def test_separate_frequency_column_supports_lai_schedule():
    csv = 'patient_id,drug,dose,unit,frequency\nP1,Zypadhera,405,mg,q4w\nP2,Okedi,75,mg,q4w\n'
    response = app.test_client().post('/upload', data={'method': 'MED', 'file': (io.BytesIO(csv.encode()), 'structured.csv')})
    wb = load_workbook(io.BytesIO(response.data))
    rows = list(wb['Results'].values)
    assert rows[1][4] == 15
    assert rows[2][4] == 11.236  # Existing oral MED table uses 0.267mg risperidone/1mg OLZ.
