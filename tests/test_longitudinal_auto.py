import csv
import io
import zipfile
from datetime import datetime

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from app import app
from services.longitudinal_auto import automatic_analysis
from tests.test_longitudinal import rx


def entries(buffer, name='Results'):
    wb = load_workbook(buffer, read_only=True)
    rows = wb[name].values
    headers = next(rows)
    output = [dict(zip(headers, row)) for row in rows]
    wb.close()
    return output


def test_automatic_reference_dates_from_rows():
    frame = pd.DataFrame([rx(), rx('2020-04-26', tabs='1')])
    results = entries(automatic_analysis({'SNU':frame}, ['DDD']))
    assert {r['기준일'] for r in results} == {'2020-04-15', '2020-04-26'}
    assert [r['결과 상태'] for r in results] == ['계산 완료', '확인 필요']
    assert results[1]['DDD (CPZ mg/day)'] is None
    assert all(r['기준일 선택'] == '각 처방일' for r in results)


def test_explicit_baseline_has_priority_per_patient():
    rows = [dict(rx(), reference_date='2020-04-20'), dict(rx('2020-04-26'), reference_date=''), dict(rx(patient='P002'), reference_date='')]
    result = entries(automatic_analysis({'Data':pd.DataFrame(rows)}, ['DDD']))
    assert [(r['patient_id'],r['기준일'],r['기준일 선택']) for r in result] == [('P001','2020-04-20','지정 기준일'),('P002','2020-04-15','각 처방일')]


def test_ambiguous_reference_dates_do_not_guess():
    with pytest.raises(ValueError, match='여러 개'):
        automatic_analysis({'Data':pd.DataFrame([dict(rx(), baseline_date='2020-04-20', mri_date='2020-04-21')])}, ['DDD'])


def test_excel_native_dates_multiple_sheets_via_main_upload():
    wb = Workbook()
    a = wb.active; a.title = 'Earlier'
    columns = list(rx())
    a.append(columns)
    a.append(list(rx(day=datetime(2020,4,15)).values()))
    b = wb.create_sheet('Later'); b.append(columns)
    b.append(list(rx(day=datetime(2020,4,26), tabs='1').values()))
    content = io.BytesIO(); wb.save(content); content.seek(0)
    response = app.test_client().post('/upload', data={'file':(content,'synthetic.xlsx'),'method':'DDD'})
    assert response.status_code == 200
    assert response.headers['Cache-Control'] == 'no-store'
    results = entries(io.BytesIO(response.data))
    assert len(results) == 2
    detail = entries(io.BytesIO(response.data), 'MedicationResults')
    assert {(r['원본 시트'], r['원본 행']) for r in detail if r['기준일'] == '2020-04-26'} == {('Earlier', 2), ('Later', 2)}


def test_mixed_static_and_longitudinal_sheets_fail_closed():
    with pytest.raises(ValueError, match='필수 열'):
        automatic_analysis({'Rx':pd.DataFrame([rx()]),'Static':pd.DataFrame([dict(patient_id='P1',drug='risperidone 2mg')])}, ['DDD'])


def test_explicit_reference_all_invalid_rx_dates_still_review():
    results = entries(automatic_analysis({'Data':pd.DataFrame([dict(rx(day='unknown'), reference_date='2020-04-20')])}, ['DDD']))
    assert results[0]['결과 상태'] == '확인 필요'


def test_main_one_line_unchanged():
    response = app.test_client().post('/api/parse', json={'text':'risperidone 2mg QD'})
    assert response.status_code == 200
    assert response.json['items'][0]['daily_dose_mg'] == 2


def test_main_csv_blank_rows_keep_original_row_numbers():
    raw = pd.DataFrame([rx(), rx('2020-04-26')]).to_csv(index=False).replace('\n', '\n\n', 1).encode()
    response = app.test_client().post('/upload', data={'file':(io.BytesIO(raw),'synthetic.csv')})
    assert response.status_code == 200
    assert {r['원본 행'] for r in entries(io.BytesIO(response.data),'MedicationResults')} == {3,4}
