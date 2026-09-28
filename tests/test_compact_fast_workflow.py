"""Regression coverage for the fast, four-sheet public workflow."""
import io
import pytest
from openpyxl import load_workbook
from app import app
from services import manual_suggestions, llm_candidates, llm_jobs
from services.converter import lookup, convert_drug
from services.result_summary import METHOD_ORDER, value_column
from tests.workbook_helpers import records

@pytest.mark.parametrize('backend', ['worker', 'ollama'])
def test_public_paths_never_search_or_wait_for_ai(monkeypatch, backend):
    monkeypatch.setenv('NAME_LLM_ENABLED', '1')
    monkeypatch.setenv('NAME_LLM_BACKEND', backend)
    def forbidden(*args, **kwargs):
        raise AssertionError('Public calculation must not search or call AI')
    monkeypatch.setattr(manual_suggestions, 'suggest_for_review', forbidden)
    monkeypatch.setattr(manual_suggestions, '_retrieval_suggestions', forbidden)
    monkeypatch.setattr(llm_candidates, 'suggest_llm', forbidden)
    monkeypatch.setattr(llm_jobs, 'online', forbidden)
    monkeypatch.setattr(llm_jobs, 'submit', forbidden)
    client = app.test_client()
    parsed = client.post('/api/parse', json={'text':'unknownxyz 2mg BID'}).json
    assert parsed['items'][0]['suggestions'] == []
    assert parsed['items'][0]['drug'] is None
    assert all(t['total_equivalent_dose_mg'] is None for t in parsed['totals'])
    assert client.get('/version').json['name_llm_enabled'] is False
    for response in [
        client.post('/api/export', json={'text':'unknownxyz 2mg BID'}),
        client.post('/upload', data={'file':(io.BytesIO(b'patient_id,medication\n001,unknownxyz 2mg BID\n'),'plain.csv')}),
        client.post('/upload', data={'file':(io.BytesIO(b'HID,PRESCR_DATE,DRUG,TABS_PER_DAY\n001,2026-09-01,unknownxyz 2mg tab,1\n'),'dated.csv')}),
    ]:
        assert response.status_code == 200
        wb = load_workbook(io.BytesIO(response.data))
        assert wb.sheetnames == ['Results','MedicationResults','Review','AuditTrail']
        assert all(s.sheet_state == 'visible' for s in wb)
        assert wb['Review'].max_row > 1
        assert records(wb, 'Results')[0]['DDD (CPZ mg/day)'] is None


def test_every_indexed_factor_matches_the_shipped_table():
    for row in lookup.itertuples(index=False):
        assert convert_drug(row.source_drug, 1.25, row.method_id, row.target_drug) == pytest.approx(1.25 * row.factor, rel=1e-14)


def test_no_duplicate_totals_or_candidate_columns_and_review_is_traceable():
    raw = b'patient_id,medication\n001,risperidone 2mg QD\n001,unknownxyz 5mg\n'
    response = app.test_client().post('/upload', data={'file':(io.BytesIO(raw),'test.csv')})
    wb = load_workbook(io.BytesIO(response.data))
    meds = records(wb, 'MedicationResults')
    assert len(meds) == 2
    assert not any('총 환산값' in h or h == 'name_candidates' for h in meds[0])
    assert meds[0][value_column('DDD')] == 120
    assert all(meds[1][value_column(m)] is None for m in METHOD_ORDER)
    issue = next(r for r in records(wb, 'Review') if 'unknownxyz' in r['original'])
    assert issue['source_row'] == 3 and issue['medication_column'] == 'medication'
    assert records(wb, 'Results')[0]['DDD (CPZ mg/day)'] is None
