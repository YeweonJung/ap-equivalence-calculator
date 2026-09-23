import json
from datetime import date
import pytest

from openpyxl import load_workbook

from services.longitudinal import analyze, analyze_export, reference_pairs
from services.result_summary import TARGETS
from tests.test_longitudinal import prepared, rx


def rows(wb, sheet):
    values = wb[sheet].values
    headers = next(values)
    return [dict(zip(headers, row)) for row in values]


def test_wide_workbook_preserves_each_method_and_date_and_review_sources():
    records = prepared([
        rx(patient='001'), rx(patient='001', drug='Risperidone 2mg tab'),
        rx(patient='001', day='2020-05-15'),
        rx(patient='002', drug='unknown 2mg tab'),
        rx(patient='', day='unknown'),
    ])
    pairs = reference_pairs(records, 'all_dates')
    expected, details = analyze(records, pairs)
    wb = load_workbook(analyze_export(records, pairs, {'policy': 'review'}))
    assert wb.sheetnames == ['Results', 'MedicationResults', 'Review']
    actual = {(r['patient_id'], r['기준일']): r for r in rows(wb, 'Results')}
    assert len(actual) == len(pairs)
    for r in expected:
        value = r['equivalent_mg']
        assert actual[(r['patient_id'], r['reference_date'])][f"{r['method']} ({TARGETS[r['method']]} mg/day)"] == (pytest.approx(value, rel=1e-12) if value is not None else None)
    exported = {(r['patient_id'], r['기준일'], r['원본 시트'] or '', r['원본 행']): r for r in rows(wb, 'MedicationResults')}
    for r in details:
        row = exported[(r['patient_id'], r['reference_date'], r['source_sheet'], r['source_row'])]
        value = r['equivalent_mg']
        assert row[f"{r['method']} ({TARGETS[r['method']]} mg/day)"] == (pytest.approx(value, rel=1e-12) if value is not None else None)
    review = rows(wb, 'Review')
    assert any(r['patient_id'] is None and '환자 ID 누락' in r['확인할 내용'] for r in review)
    assert all(wb[s].freeze_panes == 'C2' and wb[s].auto_filter.ref for s in wb.sheetnames)
    assert wb['Results']['A1'].comment and wb['Results']['A1'].fill.fgColor.rgb == 'FF1764B2'
    assert json.loads(wb.properties.description)['rule_version'] == 'prescription-date-2.0'


def test_literal_ids_and_source_text_cannot_become_formulas_or_links():
    records = prepared([rx(patient='=1+1', drug='=HYPERLINK("bad")'), rx(patient='0001'), rx(patient="'=literal")])
    wb = load_workbook(analyze_export(records, reference_pairs(records, 'all_dates'), {}))
    assert {r['patient_id'] for r in rows(wb, 'Results')} == {'=1+1', '0001', "'=literal"}
    assert {r['patient_id'] for r in rows(wb, 'MedicationResults')} == {'=1+1', '0001', "'=literal"}
    assert any(r['원문 약물'] == '=HYPERLINK("bad")' for r in rows(wb, 'MedicationResults'))
    assert not any(c.data_type == 'f' or c.hyperlink for ws in wb for row in ws for c in row)


def test_missing_factor_and_no_record_stay_blank_with_korean_guidance():
    records = prepared([rx(drug='Clopenthixol 10mg tab')])
    pairs = [('P001', date(2020, 4, 15)), ('P001', date(2021, 1, 1))]
    wb = load_workbook(analyze_export(records, pairs, {}, methods=['DDD', 'GARDNER']))
    output = rows(wb, 'Results')
    assert output[0]['DDD (CPZ mg/day)'] is None
    assert output[0]['GARDNER (CPZ mg/day)'] is not None
    assert output[1]['결과 상태'] == '해당 처방 없음'
    assert output[1]['GARDNER (CPZ mg/day)'] is None
    assert '환산 계수 없음' in rows(wb, 'Review')[0]['확인할 내용']
