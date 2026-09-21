import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from app import app
from services.manual_suggestions import suggest_for_review, SERVING_CHANNELS
from services.drug_suggestions import baseline_suggestions
from services import candidate_retrieval

ROOT = Path(__file__).resolve().parents[1]


@pytest.mark.parametrize('name,drug', [('로핀', 'zotepine'), ('조핀', 'zotepine'), ('지돈', 'ziprasidone')])
def test_app_recovers_short_candidates_only_after_user_confirmation(name, drug):
    client = app.test_client()
    result = client.post('/api/parse', json={'text': name + ' 1mg QD'}).json
    item = result['items'][0]
    assert item['status'] == 'unknown_drug'
    assert item['drug'] is None and item['conversions'] == []
    assert all(total['total_equivalent_dose_mg'] is None for total in result['totals'])
    candidate = next(c for c in item['suggestions'] if c['drug'] == drug)
    assert candidate['retrieved_by'] == ['short_hangul']
    assert candidate['confirmed_drug'] is None and candidate['auto_accepted'] is False
    assert candidate['prediction_source'] == 'edit_distance'
    assert 'short_hangul' not in candidate['explanation']
    confirmed = client.post('/api/parse', json={'text': candidate['replacement']}).json['items'][0]
    assert confirmed['drug'] == drug and confirmed['daily_dose_mg'] == 1
    assert confirmed['conversions']


@pytest.mark.parametrize('name', ['risperidonex', '로가핀', '로나핀', 'neurozapine'])
def test_unknown_and_ambiguous_app_inputs_remain_unconverted(name):
    item = app.test_client().post('/api/parse', json={'text': name + ' 1mg QD'}).json['items'][0]
    assert item['drug'] is None and item['conversions'] == []
    assert all(c['confirmed_drug'] is None and not c['auto_accepted'] for c in item['suggestions'])


def test_lr_environment_cannot_activate_model_in_serving(monkeypatch):
    from services import name_ranker
    monkeypatch.setenv('NAME_RANKER_ENABLED', '1')
    def forbidden():
        raise AssertionError('Serving must not load LR')
    monkeypatch.setattr(name_ranker, 'load_ranker', forbidden)
    assert suggest_for_review('로핀 1mg')[0]['prediction_source'] == 'edit_distance'


def test_explicit_rollback_uses_original_baseline(monkeypatch):
    monkeypatch.setenv('NAME_RETRIEVAL_ENABLED', '0')
    for text in ('로핀 1mg', 'risperidnoe 2mg BID'):
        assert suggest_for_review(text) == baseline_suggestions(text)
    assert app.test_client().get('/version').json['name_retrieval'] == 'legacy-baseline'


def test_failure_falls_back_without_logging_input(monkeypatch, caplog):
    def broken(*args, **kwargs):
        raise ValueError('sensitive exception must not appear')
    monkeypatch.setattr(candidate_retrieval, 'review_retrieval', broken)
    assert suggest_for_review('risperidnoe 2mg') == baseline_suggestions('risperidnoe 2mg')
    assert 'ValueError' in caplog.text
    assert 'sensitive' not in caplog.text and 'risperidnoe' not in caplog.text


def test_export_carries_short_candidate_evidence_without_conversion():
    response = app.test_client().post('/api/export', json={'text': '로핀 1mg QD'})
    wb = load_workbook(io.BytesIO(response.data))
    audit = dict(zip(next(wb['AuditTrail'].values), list(wb['AuditTrail'].values)[1]))
    assert audit['parsed'] is None and wb['Detailed'].max_row == 1
    candidate = json.loads(audit['name_candidates'])[0]
    assert candidate['drug'] == 'zotepine' and candidate['auto_accepted'] is False
    assert candidate['retrieval_evidence']


def test_csv_upload_uses_same_manual_retrieval():
    response = app.test_client().post('/upload', data={
        'file': (io.BytesIO('medication\n로핀 1mg QD\n'.encode('utf-8')), 'synthetic.csv'),
        'method': 'ALL'}, content_type='multipart/form-data')
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    audit = dict(zip(next(wb['AuditTrail'].values), list(wb['AuditTrail'].values)[1]))
    assert audit['parsed'] is None
    assert json.loads(audit['name_candidates'])[0]['drug'] == 'zotepine'


def test_formulation_conflict_preserves_suffix_and_requires_review():
    candidate = next(c for c in suggest_for_review('sustena oral tablet 156mg') if c['alias'] == 'sustenna')
    assert candidate['formulation_conflict'] and not candidate['auto_accepted']
    assert candidate['replacement'] == 'sustenna oral tablet 156mg'
    assert '투여경로' in candidate['explanation']


def test_serving_channels_extend_frozen_selection_with_user_requested_jamo():
    lock = json.loads((ROOT / 'data/candidate_retrieval/policy_lock.json').read_text())
    assert set(SERVING_CHANNELS) == set(lock['channels']) | {'hangul_jamo'}
    version = app.test_client().get('/version').json
    assert version['name_retrieval_channels'] == list(SERVING_CHANNELS)
    assert version['name_ranker_enabled'] is False
    assert version['automatic_confirmation_enabled'] is False


def test_app_jamo_channel_recovers_two_vowel_errors_without_confirmation():
    original = '헬로패리돌 1mg QD'
    assert baseline_suggestions(original) == []
    item = app.test_client().post('/api/parse', json={'text': original}).json['items'][0]
    assert item['drug'] is None and item['conversions'] == []
    candidate = next(c for c in item['suggestions'] if c['drug'] == 'haloperidol')
    assert candidate['retrieved_by'] == ['hangul_jamo']
    assert candidate['replacement'] == '할로페리돌 1mg QD'
    assert candidate['confirmed_drug'] is None and not candidate['auto_accepted']
    assert candidate['serving_version'] == 'candidate-retrieval-v1-jamo-manual'


def test_jamo_candidates_are_exported_without_conversion():
    response = app.test_client().post('/api/export', json={'text': '헬로패리돌 1mg QD'})
    wb = load_workbook(io.BytesIO(response.data))
    audit = dict(zip(next(wb['AuditTrail'].values), list(wb['AuditTrail'].values)[1]))
    assert audit['parsed'] is None and wb['Detailed'].max_row == 1
    candidate = json.loads(audit['name_candidates'])[0]
    assert candidate['retrieved_by'] == ['hangul_jamo']
    assert not candidate['auto_accepted']


def test_conversion_sources_and_existing_tests_remain_protected():
    from scripts.retrieval_integrity import verify_integration
    manifest = verify_integration(ROOT)
    assert set(manifest['intentional_source_changes']) == {'app.py', 'services/frames.py', 'services/release.py'}


def test_no_unresolved_merge_markers_in_runtime_sources():
    paths = [ROOT / 'app.py', ROOT / 'README.md', ROOT / 'PARSING_METHODS_KO.md']
    paths += [p for folder in ('services', 'lookup', 'templates', 'tests')
              for p in (ROOT / folder).rglob('*') if p.suffix in ('.py', '.csv', '.html')]
    markers = (chr(60) * 7 + ' ', chr(62) * 7 + ' ', chr(61) * 7)
    for path in paths:
        assert not any(line.startswith(markers) for line in path.read_text(encoding='utf-8-sig').splitlines()), path
