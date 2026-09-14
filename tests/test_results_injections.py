import io
import pandas as pd
import pytest
from openpyxl import load_workbook
from app import app
from services.result_summary import METHOD_ORDER, value_column


def parse(text):
    return app.test_client().post('/api/parse', json={'text': text}).get_json()


@pytest.mark.parametrize('text,expected', [
    ('paliperidone 100mg (PP1M)', 400),
    ('paliperidone 150mg (1개월 지속형 주사, PP1M)', 600),
    ('aripiprazole 400mg LAI q4w', 400 / 28 / 13.3 * 300),
    ('risperidone 25mg depot q2w', 25 / 14 / 2.7 * 300),
])
def test_route_specific_ddd_and_total(text, expected):
    response = parse(text)
    item = response['items'][0]
    assert item['route'] == 'injection' and item['status'] == 'converted'
    assert item['conversion_basis'] == 'WHO depot DDD'
    for c in item['conversions']:
        if c['method'] == 'DDD':
            assert c['value'] == round(expected, 4)
        elif item['drug'] != 'paliperidone':
            assert c['value'] is None
    total = next(t for t in response['totals'] if t['method'] == 'DDD')
    assert total['total_equivalent_dose_mg'] == round(expected, 4)


@pytest.mark.parametrize('text', [
    'aripiprazole 400mg LAI', 'paliperidone 100mg PP1M q4w',
    'paliperidone palmitate 156mg PP1M', 'paliperidone 100mg PP1M loading',
    'paliperidone 0mg PP1M', 'paliperidone -100mg PP1M',
    'paliperidone 100mg 150mg PP1M', 'haloperidol 50mg depot monthly',
])
def test_unverified_injection_conditions_stay_blank(text):
    response = parse(text)
    assert all(c['value'] is None for item in response['items'] for c in item['conversions'])
    assert all(t['total_equivalent_dose_mg'] is None for t in response['totals'])


def test_lai_both_never_reuses_oral_factors():
    response = parse('aripiprazole 400mg, paliperidone 150mg (LAI둘다)')
    assert len(response['items']) == 2
    assert all(i['route'] == 'injection' and i['daily_dose_mg'] is None for i in response['items'])
    assert all(t['total_equivalent_dose_mg'] is None for t in response['totals'])


def test_results_sheet_blank_failures_and_errors_for_assumed_units():
    csv = 'patient_id,medication\nP1,"ris 2, olz 5"\nP2,"ris 2mg, Unknown 3mg"\nP3,paliperidone 100mg PP1M\n'
    response = app.test_client().post('/upload', data={'method':'ALL', 'file':(io.BytesIO(csv.encode()), 'example.csv')})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames[0] == 'Results'
    sheet = wb['Results']
    headers = [c.value for c in sheet[1]]
    assert headers[3:8] == [value_column(m) for m in METHOD_ORDER]
    assert headers[8:13] == [value_column(m, True) for m in METHOD_ORDER]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert rows[0][1] == 'ris 2, olz 5'
    assert all(value is not None for value in rows[0][8:12])
    assert all(value is None for value in rows[1][8:])  # total shown once
    assert all(value is None for value in rows[2][8:])  # incomplete cell
    assert all(value is None for value in rows[3][3:])  # unknown drug
    assert rows[4][6] == 400 and rows[4][11] == 400
    assert rows[4][3] is None and rows[4][8] is None
    errors = pd.read_excel(io.BytesIO(response.data), sheet_name='Errors')
    warnings = errors.loc[errors['unit_assumed'].eq(True), 'error']
    assert len(warnings) == 2 and warnings.str.contains('확인바람').all()


def test_missing_unit_is_not_unknown_or_missing_dose():
    assert parse('olz 1.3.G BID')['items'][0]['status'] != 'converted'
    assert parse('olz')['items'][0]['status'] != 'converted'
    assert parse('olz 5 BID')['items'][0]['daily_dose_mg'] == 10


def test_structured_default_mg_stays_auditable():
    csv = 'patient_id,drug,dose,unit,frequency\nP1,"RIS,OLZ","2,5",,BID\n'
    response = app.test_client().post('/upload', data={'method':'DDD', 'file':(io.BytesIO(csv.encode()),'structured.csv')})
    sheets = pd.read_excel(io.BytesIO(response.data), sheet_name=None)
    assert sheets['Detailed']['daily_dose_mg'].tolist() == [4,10]
    assert sheets['Errors']['error'].str.contains('확인바람').all()
