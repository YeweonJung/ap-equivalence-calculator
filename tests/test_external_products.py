import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from app import app
from scripts.verify_external_products import check_items
from tests.workbook_helpers import records

DATA=json.loads((Path(__file__).resolve().parents[1]/'data/evaluation/external_products.json').read_text(encoding='utf-8'))


@pytest.mark.parametrize('case',DATA['cases'],ids=lambda case:case['id'])
def test_external_product_descriptors(case):
    response=app.test_client().post('/api/parse',json={'text':case['text']})
    assert response.status_code==200
    assert check_items(case,response.json['items'])==[]


@pytest.mark.parametrize('concentration',['1 mg/1 mL','5 mg/5ml','2 mg per 1 mL','2 mg in 1 mL','0.001 g/1ml'])
def test_oral_concentration_is_never_a_single_administered_dose(concentration):
    result=app.test_client().post('/api/parse',json={'text':f'risperidone {concentration} QD'}).json
    assert all(item['daily_dose_mg'] is None and not item['conversions'] for item in result['items'])
    assert all(total['total_equivalent_dose_mg'] is None for total in result['totals'])


@pytest.mark.parametrize('text',[
    'asenapine 3.8mg/d TRANSDERMAL QD',
    'asenapine 3.8mg patch QD',
    '아세나핀 3.8mg 패치',
    'haloperidol 5mg INTRAVENOUS 2 times a day',
])
def test_non_oral_routes_cannot_fall_back_to_oral_factors(text):
    items=app.test_client().post('/api/parse',json={'text':text}).json['items']
    assert len(items)==1
    assert items[0]['route'] in {'transdermal','injection'}
    assert items[0]['daily_dose_mg'] is None and not items[0]['conversions']
    assert items[0]['needs_review']


def test_patch_text_in_dated_upload_never_uses_tablet_factors():
    text='HID,PRESCR_DATE,DRUG,TABS_PER_DAY,COMMERCIAL_DRUG\nPATCH,2026-10-19,Asenapine 3.8mg tablet,1,Asenapine transdermal patch 3.8mg\n'
    response=app.test_client().post('/upload',data={'file':(io.BytesIO(text.encode()),'patch.csv')})
    assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.data))
    audit=records(wb,'AuditTrail')[0]
    assert audit['daily_mg'] is None
    assert 'unsupported_route' in audit['issues']
    for row in records(wb,'MedicationResults'):
        dose_columns=[key for key in row if key.endswith('mg/day)')]
        assert len(dose_columns)>1
        assert all(row[key] is None for key in dose_columns)
        assert row['상태코드']=='a' and row['검토필요'] is False
        assert '경피제·패치' in row['확인할 내용']


def test_known_oral_doses_and_validated_syringe_labels_still_convert():
    client=app.test_client()
    for text,expected in [('risperidone 2mg QD',2),('asenapine 5mg SUBLINGUAL BID',10)]:
        item=client.post('/api/parse',json={'text':text}).json['items'][0]
        assert item['daily_dose_mg']==expected and not item['needs_review']
    item=client.post('/api/parse',json={'text':'Abilify Maintena 300mg/1.5mL q4w'}).json['items'][0]
    assert item['status']=='converted' and not item['needs_review']


@pytest.mark.parametrize('concentration',['1mg/1ml','1 mg per 1 mL','1mg/1cc'])
def test_dated_product_concentration_cannot_be_used_as_tablet_strength(concentration):
    text=f'HID,PRESCR_DATE,DRUG,TABS_PER_DAY,COMMERCIAL_DRUG\nLIQUID,2026-10-19,Risperidone 1mg tablet,2,Risperidone {concentration}\n'
    response=app.test_client().post('/upload',data={'file':(io.BytesIO(text.encode()),'liquid.csv')})
    assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.data))
    audit=records(wb,'AuditTrail')[0]
    assert audit['daily_mg'] is None
    assert 'oral_solid_unconfirmed' in audit['issues']
    for row in records(wb,'MedicationResults'):
        assert all(value is None for key,value in row.items() if key.endswith('mg/day)'))


def test_concentration_guard_keeps_the_next_real_drug():
    items=app.test_client().post('/api/parse',json={'text':'risperidone 2 mg per 1 mL QD perphenazine 4mg BID'}).json['items']
    assert len(items)==2
    assert items[0]['daily_dose_mg'] is None and not items[0]['conversions']
    assert items[1]['drug']=='perphenazine' and items[1]['daily_dose_mg']==8
