import csv
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from app import app
from services.converter import lookup, convert_drug
from services.frames import parse_frames, convert_frame
from services.result_summary import METHOD_ORDER, value_column

ROOT = Path(__file__).resolve().parents[1]
CASES = list(csv.DictReader((Path(__file__).parent / 'annotated_equivalence_cases.csv').open(encoding='utf-8')))


@pytest.mark.parametrize('case', CASES, ids=lambda c: c['id'])
def test_independently_annotated_cases(case):
    frames = parse_frames(case['original'])
    assert len(frames) == 1
    item = convert_frame(frames[0], METHOD_ORDER)
    assert item['status'] == case['expected_status']
    if item['status'] != 'converted':
        assert not item['conversions']
        return
    for field, expected in [('active_moiety_mg', 'expected_active_mg'), ('oral_equivalent_mg', 'expected_oral_mg')]:
        assert item.get(field) == (float(case[expected]) if case[expected] else None)
    values = {c['method']:c['value'] for c in item['conversions']}
    for method in ['WOODS', 'GARDNER', 'DDD']:
        raw = case[f'expected_{method.lower()}_cpz']
        if raw:
            assert values[method] == pytest.approx(float(raw), abs=0.0001)
        elif method != 'DDD':
            assert values[method] is None


def test_all_factors_match_anchors_and_have_unique_positive_keys():
    assert not lookup.duplicated(['method_id', 'source_drug', 'target_drug']).any()
    assert (lookup.factor > 0).all()
    anchors = list(csv.DictReader((ROOT / 'lookup/equivalence_anchors.csv').open(encoding='utf-8')))
    for row in anchors:
        assert convert_drug(row['drug'], float(row['anchor_dose_mg']), row['method_id'], row['reference_drug']) == pytest.approx(float(row['reference_dose_mg']))
    with pytest.raises(LookupError):
        convert_drug('paliperidone', 9, 'WOODS')


def test_export_preserves_mass_audit_sources_and_separate_totals():
    content = 'patient_id,medication\nSYNTHETIC,"Sustenna 156mg, risperidone 2mg QD"\n'
    response = app.test_client().post('/upload', data={'method':'ALL', 'file':(io.BytesIO(content.encode()), 'synthetic.csv')})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    rows = [dict(zip(next(wb['Results'].values), row)) for row in list(wb['Results'].values)[1:]]
    assert rows[0][value_column('GARDNER', True)] == 800
    assert rows[0][value_column('WOODS', True)] is None
    audit = [dict(zip(next(wb['AuditTrail'].values), row)) for row in list(wb['AuditTrail'].values)[1:]]
    assert audit[0]['input_dose_mg'] == 156
    assert audit[0]['active_moiety_mg'] == 100
    assert audit[0]['oral_equivalent_mg'] == 9
    assert audit[0]['mass_source'] and audit[0]['oral_bridge_source']
    assert 'WOODS' in audit[0]['unavailable_methods']
    assert wb['FactorSources'].max_row > 40
    assert 'ReviewQueue' in wb.sheetnames and 'VersionInfo' in wb.sheetnames


@pytest.mark.parametrize('text', ['Sustenna 156mg/mL monthly', 'Aristada 441mg INITIO',
                                 'Sustenna 156mg active moiety', 'Hafyera 1560mg PP1M'])
def test_unsafe_mass_or_product_input_is_not_converted(text):
    items = [convert_frame(f, METHOD_ORDER) for f in parse_frames(text)]
    assert all(i['status'] != 'converted' for i in items)
