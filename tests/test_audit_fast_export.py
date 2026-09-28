import io
import re
import pandas as pd
import pytest
from openpyxl import load_workbook
from app import app
from services.frames import alias_map, drug_mentions, _outside
from services.longitudinal_auto import automatic_analysis
from tests.workbook_helpers import records
from tests.test_longitudinal import rx


def test_audit_restores_mass_assumptions_and_unconverted_records():
    raw = 'patient_id,medication\n001,"Sustenna 156mg, risperidone 2mg QD, unknownxyz 5mg"\n002,lithium 600mg\n'
    response = app.test_client().post('/upload', data={'file':(io.BytesIO(raw.encode()),'input.csv')})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames == ['Results','MedicationResults','Review','AuditTrail']
    audit = records(wb, 'AuditTrail')
    assert len(audit) == 4
    assert [r['status'] for r in audit] == ['converted','converted','unknown_drug','non_target']
    injection = audit[0]
    assert injection['input_dose_mg'] == 156 and injection['active_moiety_mg'] == 100
    assert injection['oral_equivalent_mg'] == 9 and injection['interval_days'] == 30
    assert injection['mass_source'] and injection['oral_bridge_source']
    assert 'WOODS' in injection['unavailable_methods']
    assert injection['source_row'] == 2 and injection['patient'] == '001'
    for row in audit[:3]:
        assert row['source_text'][row['source_start']:row['source_end']] == row['original']
    assert audit[2]['parsed'] is None and audit[2]['daily_dose_mg'] is None
    assert audit[2]['name_candidates'] == '[]'
    assert wb['AuditTrail'].sheet_state == 'visible'
    assert wb['AuditTrail'].auto_filter.ref
    assert wb['AuditTrail'].freeze_panes == 'E2'


def test_dated_audit_includes_invalid_and_excluded_source_rows():
    data = [rx(patient='001', drug='Risperidone 2mg tab', tabs='1'),
            rx(patient='001', drug='Lithium 600mg tab', tabs='1'),
            rx(patient='', day='not-a-date', drug='unknownxyz 2mg tab')]
    wb = load_workbook(automatic_analysis({'Source':pd.DataFrame(data)}, ['DDD']))
    audit = records(wb, 'AuditTrail')
    assert len(audit) == 3
    assert [r['source_row'] for r in audit] == [2,3,4]
    assert all(r['source_sheet'] == 'Source' for r in audit)
    assert audit[0]['daily_mg'] == 2
    assert audit[1]['kind'] == 'non_target'
    assert audit[2]['prescription_date'] == 'not-a-date'
    assert 'missing_patient' in audit[2]['issues'] and 'invalid_date' in audit[2]['issues']
    assert all(r['original_end'] is None and r['effective_end'] is None for r in audit)


def test_audit_preserves_structured_source_and_literal_ids():
    raw = 'patient_id,drug,dose,unit,frequency\n=1+1,risperidone,0.001,mg,QD\n'
    response = app.test_client().post('/upload', data={'file':(io.BytesIO(raw.encode()),'input.csv')})
    wb = load_workbook(io.BytesIO(response.data))
    audit = records(wb, 'AuditTrail')[0]
    assert audit['patient'] == '=1+1'
    assert audit['source_text'] == 'risperidone'
    assert audit['original'] == 'risperidone 0.001mg QD'
    assert audit['dose_mg'] == pytest.approx(0.001)
    assert not any(c.data_type == 'f' or c.hyperlink for ws in wb for row in ws for c in row)


def test_precompiled_mentions_keep_all_alias_boundaries_and_overlap_order():
    patterns = [(re.compile(r'(?<![\w])' + re.escape(alias) + r'(?![a-z가-힣])', re.I), name)
                for alias, name in alias_map.items()]
    def reference(text):
        hits = []
        masked = _outside(text)
        for pattern, name in patterns:
            hits.extend((m.start(),m.end(),name) for m in pattern.finditer(masked))
        chosen = []
        for hit in sorted(hits,key=lambda h:(h[0],-(h[1]-h[0]))):
            if not chosen or hit[0] >= chosen[-1][1]:chosen.append(hit)
        return chosen
    for alias in alias_map:
        for text in (alias.upper()+'2mg QD; risperidone 1mg', 'x'+alias+'x 2mg (olanzapine)'):
            assert drug_mentions(text) == reference(text)
