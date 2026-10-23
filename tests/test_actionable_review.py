"""Only demonstrated errors enter Review; uncertainty still blocks full totals."""
import io
import pandas as pd
import pytest
from openpyxl import load_workbook
from app import app
from services.longitudinal_auto import automatic_analysis
from tests.workbook_helpers import records


def upload(text, method='ALL'):
    raw = 'patient_id,medication\nP001,"' + text + '"\n'
    response = app.test_client().post('/upload', data={
        'method': method, 'file': (io.BytesIO(raw.encode()), 'synthetic.csv')})
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.data))


@pytest.mark.parametrize('text,code', [
    ('unknownxyz 5mg QD', 'b'), ('risperidone 2 QD', 'c'),
    ('risperidone 2mg', 'c'), ('aripiprazole 400mg LAI', 'c'),
    ('blonanserin 8mg QD', 'a'), ('asenapine 3.8mg patch', 'a'),
    ('lorazepam 1mg QHS', 'e'),
])
def test_information_codes_stay_outside_review_in_every_sheet(text, code):
    wb = upload(text, 'DDD')
    assert records(wb, 'Review') == []
    for sheet in ('보완합계', '보완근거', 'Results', 'MedicationResults', 'AuditTrail'):
        row = records(wb, sheet)[0]
        assert code in row['상태코드'].split(';'), (sheet, row)
        assert row['검토필요'] is False
    assert len(records(wb, '코드안내')) == 7


def test_partial_method_coverage_preserves_nulls_and_numbers():
    wb = upload('risperidone 2mg QD')
    assert records(wb, 'Review') == []
    assert records(wb, 'MedicationResults')[0]['상태코드'] == 'a'
    result = records(wb, 'Results')[0]
    assert result['DDD (OLZ mg/day)'] == 4
    assert result['CPZ_FGA (CPZ mg/day)'] is None


def test_companion_does_not_hide_unknown_or_allow_incomplete_total():
    wb = upload('lorazepam 1mg QHS, unknownxyz 5mg QD, risperidone 2mg QD', 'DDD')
    assert records(wb, 'Review') == []
    assert len(records(wb, 'AuditTrail')) == 3
    assert records(wb, 'Results')[0]['DDD (OLZ mg/day)'] is None
    summary = records(wb, '보완합계')[0]
    assert summary['보완 총합 (mg/day)'] is None
    assert summary['부분합 (mg/day)'] == 4
    assert summary['미환산 수'] == 1


@pytest.mark.parametrize('text', ['risperidone -2mg QD', 'risperidone 0mg QD',
                                  'risperidone 2mg QD BID', 'paliperidone 150mg LAI q0w'])
def test_actual_errors_only_are_review_required(text):
    wb = upload(text, 'DDD')
    review = records(wb, 'Review')
    assert len(review) == 1 and review[0]['검토필요'] is True
    assert 'x' in review[0]['상태코드']
    assert records(wb, 'Results')[0]['DDD (OLZ mg/day)'] is None
    item = app.test_client().post('/api/parse', json={'text': text}).json['items'][0]
    assert item['review_required'] and not item['ok']


@pytest.mark.parametrize('methods', [['DDD','CPZ_FGA'], ['CPZ_FGA']])
def test_dated_missing_coefficients_are_not_errors(methods):
    df = pd.DataFrame({'patient_id':['001'],'prescription_date':['2026-10-23'],
                       'drug_name':['risperidone 2mg tab'],'tablets_per_day':['1']})
    wb = load_workbook(automatic_analysis({'input':df}, methods))
    assert records(wb, 'Review') == []
    detail = records(wb, 'MedicationResults')[0]
    assert detail['상태코드'] == 'a'
    assert detail['CPZ_FGA (CPZ mg/day)'] is None
    if 'DDD' in methods:
        assert detail['DDD (OLZ mg/day)'] == 4


def test_dated_invalid_values_distinguished_from_missing_and_unknown():
    df = pd.DataFrame({'patient_id':['unknown','missing','negative','date','conflict'],
                      'prescription_date':['2026-10-23'] * 3 + ['2026-13-99', '2026-10-23'],
                      'drug_name':['unknownxyz 5mg tab'] + ['risperidone 2mg tab'] * 4,
                      'product_name':['', '', '', '', 'olanzapine 2mg tab'],
                      'tablets_per_day':['1', '', '-1', '1', '1']})
    wb = load_workbook(automatic_analysis({'input':df}, ['DDD']))
    review = records(wb, 'Review')
    assert {r['patient_id'] for r in review} == {'negative', 'date', 'conflict'}
    audit = {r['patient_id']: r for r in records(wb, 'AuditTrail')}
    assert audit['unknown']['상태코드'] == 'b'
    assert audit['missing']['상태코드'] == 'c'
    assert all(r['검토필요'] for r in review)


def test_duplicate_remains_a_blocker_without_becoming_a_confirmed_error():
    df = pd.DataFrame({'patient_id':['001','001'], 'prescription_date':['2026-10-23'] * 2,
                       'drug_name':['risperidone 2mg tab'] * 2,'tablets_per_day':['1'] * 2})
    wb = load_workbook(automatic_analysis({'input':df}, ['DDD']))
    assert not records(wb, 'Review')
    assert 'd' in records(wb, 'AuditTrail')[1]['상태코드']
    assert records(wb, 'Results')[0]['DDD (OLZ mg/day)'] is None


def test_homepage_explains_codes_and_unknown_is_not_an_error():
    client = app.test_client()
    html = client.get('/').get_data(as_text=True)
    assert '결과 상태 코드 안내' in html and 'Review 시트에는 x 오류만' in html
    data = client.post('/api/parse', json={'text':'unknownxyz 5mg QD'}).json
    assert data['items'][0]['status_code'] == 'b'
    assert data['items'][0]['review_required'] is False
    assert data['items'][0]['needs_review'] is True
    assert data['cases'][0]['total'] is None
