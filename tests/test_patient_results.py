import io
import pytest
from openpyxl import load_workbook, Workbook
from app import app
from services.result_summary import TARGETS


def upload(content):
    response = app.test_client().post('/upload', data={'file': (io.BytesIO(content.encode()), 'patients.csv')})
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.data))


def rows(wb, sheet):
    values = list(wb[sheet].values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def test_patient_sums_across_rows_keep_ids_and_method_order():
    wb = upload('patient_id,medication\n001,risperidone 2mg QD\n001,olanzapine 5mg QD\n002,risperidone 1mg QD\nNA,risperidone 1mg QD\n')
    result = rows(wb, 'Results')
    assert [r['patient_id'] for r in result] == ['001', '002', 'NA']
    assert list(result[0])[:7] == ['patient_id', 'CMD (CPZ mg/day)', 'MED (OLZ mg/day)', 'DDD (CPZ mg/day)', 'ED95 (OLZ mg/day)', 'GARDNER (CPZ mg/day)', 'WOODS (CPZ mg/day)']
    expected = app.test_client().post('/api/parse', json={'text': 'risperidone 2mg QD; olanzapine 5mg QD'}).json['totals']
    for total in expected:
        column = f"{total['method']} ({TARGETS[total['method']]} mg/day)"
        assert result[0][column] == total['total_equivalent_dose_mg']


def test_incomplete_patient_never_exports_partial_as_total():
    wb = upload('patient_id,medication\nP001,risperidone 2mg QD\nP001,unknownxyz 5mg\nP002,risperidone 2mg QD\n')
    result = rows(wb, 'Results')
    assert all(v is None for k,v in result[0].items() if k != 'patient_id')
    assert result[1]['DDD (CPZ mg/day)'] is not None
    check = next(r for r in rows(wb, 'PatientChecks') if r['patient_id']=='P001' and r['method']=='DDD')
    assert check['partial_equivalent_dose_mg'] > 0
    assert check['unresolved_count'] == 1


def test_missing_ids_are_separate_and_empty_prescription_is_blank():
    wb = upload('patient_id,medication\n,risperidone 2mg QD\n,olanzapine 5mg QD\nP003,\n')
    result = rows(wb, 'Results')
    assert len(result) == 3
    assert len({r['patient_id'] for r in result}) == 3
    assert result[2]['DDD (CPZ mg/day)'] is None
    assert len(rows(wb, 'Errors')) == 2


def test_excel_multisheet_and_multiple_medication_columns():
    wb = Workbook()
    for ws in [wb.active, wb.create_sheet('second')]:
        ws.append(['patient_id', 'medication1', 'medication2'])
        ws.append(['001', 'risperidone 2mg QD', 'olanzapine 5mg QD'])
    content=io.BytesIO(); wb.save(content); content.seek(0)
    response=app.test_client().post('/upload', data={'file':(content,'multi.xlsx')})
    actual=load_workbook(io.BytesIO(response.data))
    result=rows(actual,'Results')
    assert len(result)==1 and result[0]['patient_id']=='001'
    totals=app.test_client().post('/api/parse',json={'text':'risperidone 4mg QD; olanzapine 10mg QD'}).json['totals']
    expected=next(t['total_equivalent_dose_mg'] for t in totals if t['method']=='DDD')
    assert result[0]['DDD (CPZ mg/day)']==expected


def test_formula_like_id_is_exact_literal_text():
    wb = upload('patient_id,medication\n=1+1,risperidone 2mg QD\n')
    assert wb['Results']['A2'].value == '=1+1'
    assert wb['Results']['A2'].data_type == 's'
