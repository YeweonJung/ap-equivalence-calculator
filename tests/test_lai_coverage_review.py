import io
import json
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from app import app
from services.frames import parse_frames, convert_frame
from services.longitudinal import prepare, detect_mapping, analyze, analyze_export
from tests.workbook_helpers import records


@pytest.mark.parametrize('text,oral,ddd', [
    ('Longavo 50mg monthly', 2, 50/30/2.7*10),
    ('Longavo 125mg SC monthly', 5, 125/30/2.7*10),
    ('Rykindo 37.5mg q2w', None, 37.5/14/2.7*10),
    ('Risvan 75mg monthly', 3, 75/30/2.7*10),
    ('Risvan 100mg monthly', 4, 100/30/2.7*10),
    ('Haldol decanoate 100mg q4w', None, 100/28/3.3*10),
    ('할돌 데카노아스 50mg 4주', None, 50/28/3.3*10),
    ('Palmeux 75mg', 6, 75/30/2.5*10),
    ('Paliperidone Teva injection 100mg monthly', 9, 100/30/2.5*10),
    ('Niapelf 150mg', 12, 150/30/2.5*10),
])
def test_new_products_have_verified_values_and_separate_notes(text, oral, ddd):
    items = parse_frames(text)
    assert len(items) == 1
    item = convert_frame(items[0], ['DDD', 'MED'])
    assert item['status'] == 'converted', item
    assert item['oral_equivalent_mg'] == oral
    assert item['conversions'][0]['value'] == pytest.approx(ddd, abs=0.0001)
    assert not item['needs_review'] and not item['warning']
    assert item['calculation_note'] and item['product_name_checked_on'] == '2026-10-15'
    if oral is None:
        assert item['conversions'][1]['value'] is None


@pytest.mark.parametrize('text', [
    'Longavo 100mg q2mo', 'Longavo 100mg IM monthly', 'Longavo 100mg',
    'Rykindo 12.5mg q2w', 'Rykindo 25mg q4w', 'Rykindo Consta 25mg q2w',
    'Risvan 100mg q4w', 'Haldol decanoate 400mg q4w',
    'Haldol decanoate 100mg', 'Haldol decanoate 50mg/ml q4w',
    'Haldol decanoate 100mg IV q4w', 'Haldol decanoate 70.52mg ester q4w',
    'Niapelf 100mg loading', 'Niapelf 156mg', 'Palmeux 100mg q4w',
    'Maintena 400mg 4주 6주', 'Maintena 400mg 1.5주',
])
def test_real_ambiguity_still_blocks_values(text):
    items = [convert_frame(f, ['DDD', 'MED']) for f in parse_frames(text)]
    assert all(i['needs_review'] for i in items)
    assert not any(c['value'] is not None for i in items for c in i['conversions'])


@pytest.mark.parametrize('schedule', ['4주', '4주마다', '4주 간격', '28일마다', '28일 간격', 'q4w'])
def test_explicit_korean_intervals(schedule):
    f = parse_frames('메인테나 400mg ' + schedule)[0]
    assert f['status'] == 'ready' and f['interval_days'] == 28


def test_informational_notes_are_kept_out_of_review_but_exported():
    raw = 'patient_id,medication\nP001,Longavo 75mg monthly\n'
    r = app.test_client().post('/upload', data={'method': 'MED', 'file': (io.BytesIO(raw.encode()), 'synthetic.csv')})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert not records(wb, 'Review')
    assert records(wb, 'AuditTrail')[0]['calculation_note']
    assert '경구 대응용량 기반 추정' in records(wb, 'MedicationResults')[0]['환산 근거']
    f = parse_frames('Maintena 400 LAI q4w')[0]
    assert f['needs_review'] and f['unit_assumed']


def dated(drug, product='', daily='', days='999', patient='001'):
    df = pd.DataFrame([dict(HID=patient, PRESCR_DATE='2026-10-15', PRESCR_DAYS=days,
                           TABS_PER_DAY=daily, DRUG=drug, COMMERCIAL_DRUG=product)])
    return prepare(df, detect_mapping(df.columns), 'prescription_date')


def test_dated_lai_uses_interval_not_days_supply_and_keeps_bridge_in_export():
    source = dated('risperidone 75mg', 'Longavo 75mg monthly')
    assert not source[0]['issues']
    pairs = [('001', date(2026, 10, 15))]
    results, detail = analyze(source, pairs, ['DDD', 'MED'], 'prescription_date')
    assert results[0]['equivalent_mg'] == pytest.approx(75/30/2.7*10)
    assert results[1]['equivalent_mg'] == pytest.approx(3/0.267)
    assert json.loads(detail[0]['injection_evidence'])['route'] == 'injection'
    wb = load_workbook(analyze_export(source, pairs, {'policy': 'prescription_date'},
                                     ['DDD', 'MED'], 'prescription_date'))
    assert not records(wb, 'Review')
    evidence = records(wb, 'AuditTrail')[0]['injection_evidence']
    assert json.loads(evidence)['interval_days'] == 30
    # Supplemental cases must use the LAI frame, never treat average injected mg as oral mg.
    from services.supplemented import dated_cases
    stream = io.BytesIO()
    import csv
    text = io.TextIOWrapper(stream, encoding='utf-8-sig', newline='')
    writer = csv.DictWriter(text, fieldnames=list(detail[0]))
    writer.writeheader(); writer.writerows(detail); text.flush()
    cases = list(dated_cases(stream, results, base_method='MED'))
    assert next(c for c in cases if c['base'] == 'MED')['total'] == pytest.approx(round(3/0.267, 4))
    text.detach()


@pytest.mark.parametrize('drug,product,daily', [
    ('risperidone 75mg', 'Longavo 100mg monthly', ''),
    ('olanzapine 75mg', 'Longavo 75mg monthly', ''),
    ('risperidone 75mg', 'Longavo 75mg', ''),
    ('risperidone 75mg', 'Longavo 75mg monthly', '2'),
    ('risperidone 75mg QD', 'Longavo 75mg monthly', ''),
])
def test_dated_ambiguity_is_not_cleared(drug, product, daily):
    source = dated(drug, product, daily)
    results, _ = analyze(source, [('001', date(2026, 10, 15))], ['DDD'], 'prescription_date')
    assert source[0]['issues'] and results[0]['equivalent_mg'] is None


def test_dated_equivalent_paliperidone_strengths_and_cache_separation():
    source = dated('paliperidone 100mg', 'Invega Sustenna 156mg')
    assert not source[0]['issues']
    assert source[0]['injection_frame']['active_moiety_mg'] == 100
    same_label = dated('paliperidone 156mg', 'Invega Sustenna 156mg')
    assert not same_label[0]['issues']
    assert same_label[0]['injection_frame']['active_moiety_mg'] == 100
    # Same mass/day across oral and injection must not collide in the conversion cache.
    a = dated('risperidone 75mg', 'Longavo 75mg monthly')
    b = dated('risperidone 2.5mg tab', daily='1', patient='002')
    out, _ = analyze(a+b, [('001', date(2026,10,15)), ('002', date(2026,10,15))], ['MED'], 'prescription_date')
    assert out[0]['equivalent_mg'] == pytest.approx(3/0.267)
    assert out[1]['equivalent_mg'] == pytest.approx(2.5/0.267)
