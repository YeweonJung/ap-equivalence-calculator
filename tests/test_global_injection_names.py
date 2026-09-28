import io
import pytest
from openpyxl import load_workbook
from app import app
from services.lai_support import NAME_CATALOG, NAME_EVIDENCE
from tests.workbook_helpers import records


@pytest.mark.parametrize('alias,record', [(a,r) for a,r in NAME_EVIDENCE.items() if r['policy']=='review'])
def test_review_only_names_never_reuse_oral_or_other_product_factors(alias, record):
    data = app.test_client().post('/api/parse', json={'text': alias+' 25mg q4w'}).get_json()
    assert len(data['items']) == 1
    item = data['items'][0]
    assert item['drug'] == record['drug']
    assert item['route'] == 'injection'
    assert item['status'] == 'unsupported_formulation'
    assert item['product_name_source'] == record['source']
    assert all(c['value'] is None for c in item['conversions'])


@pytest.mark.parametrize('text,profile,oral', [
    ('Zyprexa-Relprevv 300mg q4w', 'OLZ_PAMOATE', 10),
    ('Paliperidone Zentiva injection 100mg monthly', 'PP1M', 9),
    ('인베가하피에라주사1092mg', 'PP6M', 9),
    ('유제디주100mg q2mo', 'RIS_UZEDY', 2),
])
def test_verified_new_spelling_uses_existing_validated_mapping(text, profile, oral):
    items = app.test_client().post('/api/parse', json={'text':text}).get_json()['items']
    assert len(items)==1
    assert items[0]['status']=='converted'
    assert items[0]['lai_profile']==profile
    assert items[0]['oral_equivalent_mg']==oral
    assert items[0]['product_name_source'].startswith('https://')


@pytest.mark.parametrize('text', ['유제디주100mg', '유제디원엠주100mg q2mo',
    'Zyprexa intramuscular 300mg q4w', 'AristadaInitio 675mg q4w',
    'Rykindo Consta 25mg q2w', 'Zyprexa Relprevv 300mg SC q4w',
    'Zyprexa Relprevv 300mg oral q4w', '유제디주100mg 경구 q2mo'])
def test_missing_interval_and_conflicting_formulations_stay_blank(text):
    items=app.test_client().post('/api/parse',json={'text':text}).get_json()['items']
    assert all(c['value'] is None for i in items for c in i['conversions'])


@pytest.mark.parametrize('text', ['zyprexa 10mg QD','geodon 20mg QD','paliperidone 6mg QD'])
def test_oral_names_are_not_reclassified_as_injections(text):
    item=app.test_client().post('/api/parse',json={'text':text}).get_json()['items'][0]
    assert item['route']=='oral' and item['status']=='converted'


def test_catalog_has_sources_and_export_preserves_name_evidence():
    assert all(r['source'].startswith('https://') and r['checked_on']=='2026-09-28' for r in NAME_CATALOG)
    raw='patient_id,medication\nDEMO001,Rykindo 25mg q2w\nDEMO002,유제디주100mg q2mo\n'
    response=app.test_client().post('/upload',data={'file':(io.BytesIO(raw.encode()),'example.csv')})
    assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames==['Results','MedicationResults','Review','AuditTrail']
    audit=records(wb,'AuditTrail')
    assert audit[0]['recognized_product']=='Rykindo'
    assert audit[0]['product_name_source'].startswith('https://dailymed.nlm.nih.gov/')
    assert audit[0]['status']=='unsupported_formulation'
    assert audit[1]['original']=='유제디주100mg q2mo'
    assert audit[1]['oral_equivalent_mg']==2
