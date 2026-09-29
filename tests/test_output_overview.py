import io
from openpyxl import load_workbook
from app import app


def test_overview_keeps_values_and_explains_blanks():
    raw='patient_id,medication\n001,risperidone 2mg QD\n002,blonanserin 8mg\n'
    response=app.test_client().post('/upload',data={'file':(io.BytesIO(raw.encode()),'test.csv')})
    wb=load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames[0]=='Results'
    sheet=wb.active
    assert sheet['A2'].value=='001' and sheet['A2'].data_type=='s'
    headers=[c.value for c in sheet[1]]
    ddd=headers.index('DDD\nOLZ mg/day')+1
    assert sheet.cell(2,ddd).value==4
    assert sheet.cell(3,ddd).value=='—'
    assert sheet['B3'].value=='계산 보류'
    assert '환산계수 없음' in sheet.cell(3,len(headers)).value
    assert wb.sheetnames == ['Results','MedicationResults','Review','AuditTrail']
    assert all(s.sheet_state == 'visible' for s in wb)
    assert wb['Results'].sheet_state=='visible'
    assert sheet.freeze_panes is None and sheet.auto_filter.ref=='A1:L3'
    wb.close()


def test_dated_overview_does_not_merge_dates():
    raw='HID,PRESCR_DATE,DRUG,TABS_PER_DAY\n001,2026-09-01,Risperidone 2mg tab,1\n001,2026-09-15,Risperidone 2mg tab,2\n'
    response=app.test_client().post('/upload',data={'file':(io.BytesIO(raw.encode()),'test.csv')})
    wb=load_workbook(io.BytesIO(response.data))
    s=wb.active
    assert s.title=='Results'
    assert s['B2'].value=='2026-09-01' and s['B3'].value=='2026-09-15'
    headers=[c.value for c in s[1]]
    ddd=headers.index('DDD\nOLZ mg/day')+1
    assert s.cell(2,ddd).value==4 and s.cell(3,ddd).value==8
    assert s.freeze_panes is None
    wb.close()
