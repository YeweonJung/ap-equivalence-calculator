import io
from copy import copy
from datetime import date

import pytest
from openpyxl import load_workbook

from app import app
from services import supplemented
from services.longitudinal import analyze_export, reference_pairs
from tests.test_longitudinal import prepared, rx

OLD_SHEETS = ['Results', 'MedicationResults', 'Review', 'AuditTrail']
ALL_SHEETS = OLD_SHEETS + ['보완합계', '보완근거']


def rows(wb, sheet):
    data = list(wb[sheet].values)
    return [dict(zip(data[0], row)) for row in data[1:]]


def upload(source, method='ALL'):
    response = app.test_client().post('/upload', data={
        'file': (io.BytesIO(source.encode()), 'synthetic.csv'), 'method': method})
    assert response.status_code == 200
    return load_workbook(io.BytesIO(response.data))


@pytest.mark.parametrize('source', [
    'patient_id,medication\n001,olanzapine 5mg QD\n001,paliperidone 6mg QD\n002,unknownxyz 3mg\n003,lithium 600mg\n004,Sustenna 156mg\n',
    'HID,PRESCR_DATE,DRUG,TABS_PER_DAY\n001,2026-10-01,Olanzapine 5mg tab,1\n001,2026-10-01,Paliperidone 6mg tab,1\n002,2026-10-01,unknownxyz 3mg tab,1\n003,2026-10-01,Lithium 600mg tab,1\n',
])
def test_supplemented_export_preserves_all_original_cells_and_formatting(source, monkeypatch):
    updated = upload(source)
    with monkeypatch.context() as patch:
        patch.setattr(supplemented, 'append_frame_sheets', lambda *args: None)
        patch.setattr(supplemented, 'append_dated_sheets', lambda *args: None)
        original = upload(source)
    assert updated.sheetnames == ALL_SHEETS
    assert original.sheetnames == OLD_SHEETS
    for name in OLD_SHEETS:
        def cells(wb):
            return [[(c.value, c.data_type, c.number_format, copy(c.alignment)) for c in row] for row in wb[name]]
        assert cells(updated) == cells(original)
    summary = rows(updated, '보완합계')
    cmd = next(r for r in summary if r['환자 ID'] == '001' and r['기본 방법'] == 'CMD')
    assert cmd['보완 총합 (mg/day)'] == 20
    assert cmd['보완 약물 수'] == 1 and cmd['미환산 수'] == 0
    assert cmd['사용 방법'] == 'CMD + MED'
    unknown = [r for r in summary if r['환자 ID'] == '002']
    assert all(r['보완 총합 (mg/day)'] is None and r['미환산 수'] == 1 for r in unknown)
    excluded = [r for r in summary if r['환자 ID'] == '003']
    assert all(r['보완 총합 (mg/day)'] == 0 for r in excluded)
    for name in ALL_SHEETS:
        ws = updated[name]
        assert ws.freeze_panes is None and not ws.merged_cells.ranges
        assert ws.auto_filter.ref.startswith('A1:')
    detail = next(r for r in rows(updated, '보완근거') if r['환자 ID'] == '001' and r['기본 방법'] == 'CMD' and r['성분'] == 'paliperidone')
    assert detail['적용 방법'] == 'MED' and detail['기준약물'] == 'OLZ'
    assert detail['계산 용량 (mg/day)'] == 6 and detail['적용계수'] == 2.5
    assert detail['환산값 (mg/day)'] == 15
    assert '24493852' in detail['출처'] and detail['원본 행'] == 3


def test_selected_method_still_can_use_other_methods_and_partial_is_not_total():
    wb = upload('patient_id,medication\n0001,paliperidone 6mg QD\n0001,unknownxyz 2mg\n', 'CMD')
    result, = rows(wb, '보완합계')
    assert result['환자 ID'] == '0001'
    assert result['기본 방법'] == 'CMD'
    assert result['보완 총합 (mg/day)'] is None and result['부분합 (mg/day)'] == 15
    assert result['미환산 수'] == 1 and result['보완 약물 수'] == 1


@pytest.mark.parametrize('problem', ['duplicate', 'ingredient', 'dose', 'invalid_date'])
def test_dated_blockers_cannot_be_repaired_by_fallback(problem):
    data = [rx(patient='001', drug='Paliperidone 6mg tab', tabs='1')]
    if problem == 'duplicate':
        data.append(dict(data[0]))
    elif problem == 'ingredient':
        data[0]['COMMERCIAL_DRUG'] = 'Olanzapine 6mg tab'
    elif problem == 'dose':
        data[0]['TABS_PER_DAY'] = 'bad'
    else:
        data.append(rx(patient='001', day='bad', drug='Paliperidone 6mg tab', tabs='1'))
    records = prepared(data, 'prescription_date')
    wb = load_workbook(analyze_export(records, [('001', date(2020, 4, 15))],
                                    {'policy': 'prescription_date'}, methods=['CMD'], policy='prescription_date'))
    summary, = rows(wb, '보완합계')
    assert summary['보완 총합 (mg/day)'] is None and summary['미환산 수'] > 0


def test_reference_dates_and_no_record_are_preserved():
    records = prepared([rx(drug='Paliperidone 6mg tab', tabs='1')])
    pairs = [('P001', date(2020, 4, 15)), ('P001', date(2021, 1, 1))]
    wb = load_workbook(analyze_export(records, pairs, {}, methods=['CMD']))
    first, empty = rows(wb, '보완합계')
    assert first['보완 총합 (mg/day)'] == 15
    assert empty['처리'] == '처방 없음' and empty['보완 총합 (mg/day)'] is None


def test_same_patient_dates_not_combined_and_literal_ids():
    wb = upload('HID,PRESCR_DATE,DRUG,TABS_PER_DAY\n=1+1,2026-10-01,Paliperidone 6mg tab,1\n=1+1,2026-10-02,Olanzapine 5mg tab,1\n', 'CMD')
    first, second = rows(wb, '보완합계')
    assert first['환자 ID'] == second['환자 ID'] == '=1+1'
    assert first['보완 총합 (mg/day)'] == 15 and second['보완 총합 (mg/day)'] == 5
    assert not any(c.data_type == 'f' or c.hyperlink for ws in wb for row in ws for c in row)


def test_api_export_and_injection_provenance_arithmetic():
    response = app.test_client().post('/api/export', json={'text': 'Sustenna 156mg'})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data))
    assert wb.sheetnames == ALL_SHEETS
    detail = rows(wb, '보완근거')
    cmd = next(r for r in detail if r['기본 방법'] == 'CMD')
    assert cmd['적용 방법'] == 'MED' and cmd['계산 용량 (mg/day)'] == 9
    assert cmd['환산값 (mg/day)'] == 22.5
    assert '제품별' in cmd['용량 기준']
    assert len(cmd['출처'].split(';')) > 1
    for row in detail:
        if row['적용계수'] is not None:
            assert row['계산 용량 (mg/day)'] * row['적용계수'] == pytest.approx(row['환산값 (mg/day)'])
