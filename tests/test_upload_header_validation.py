import io
import pytest
from openpyxl import Workbook, load_workbook
from app import app

HEADER = 'HID,PRESCR_DATE,DRUG,TABS_PER_DAY'
ROW = '001,2026-09-01,Quetiapine 25mg tab,1'


@pytest.mark.parametrize('route', ['/upload', '/api/longitudinal/analyze'])
def test_leading_blank_lines_preserve_date_group_and_source_row(route):
    raw = ('\n\n' + HEADER + '\n' + ROW + '\n').encode()
    response = app.test_client().post(route, data={'file':(io.BytesIO(raw),'test.csv'),'ack':'yes'})
    assert response.status_code == 200
    wb = load_workbook(io.BytesIO(response.data),read_only=True)
    vals = wb['MedicationResults'].values
    header = next(vals)
    rows = [dict(zip(header,r)) for r in vals]
    assert len(rows) == 1
    assert rows[0]['patient_id'] == '001'
    assert rows[0]['처방일'] == '2026-09-01'
    assert rows[0]['원본 행'] == 4
    wb.close()


@pytest.mark.parametrize('route', ['/upload','/api/longitudinal/inspect','/api/longitudinal/analyze'])
@pytest.mark.parametrize('prefix', ['', '\n'])
def test_duplicate_date_headers_are_rejected_in_csv(route,prefix):
    raw = (prefix+'HID,PRESCR_DATE,PRESCR_DATE,DRUG,TABS_PER_DAY\n001,2026-09-01,2026-09-15,Quetiapine 25mg tab,1\n').encode()
    response = app.test_client().post(route,data={'file':(io.BytesIO(raw),'test.csv'),'ack':'yes'})
    assert response.status_code == 400
    assert '중복' in response.get_data(as_text=True) or '중복' in response.json['error']


@pytest.mark.parametrize('title_rows',[0,1])
def test_duplicate_headers_in_excel_cannot_be_silently_renamed(title_rows):
    wb=Workbook(); ws=wb.active
    if title_rows: ws.append(['Prescription export'])
    ws.append(['HID','PRESCR_DATE','PRESCR_DATE','DRUG','TABS_PER_DAY'])
    ws.append(['001','2026-09-01','2026-09-15','Quetiapine 25mg tab','1'])
    buf=io.BytesIO();wb.save(buf);buf.seek(0)
    response=app.test_client().post('/upload',data={'file':(buf,'test.xlsx')})
    assert response.status_code==400
    assert '중복' in response.get_data(as_text=True)
