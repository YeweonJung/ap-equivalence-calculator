from services.exclusions import is_excluded, exclusion_label
"""Compact presentation layer; underlying numeric results stay unchanged."""
from collections import defaultdict
from services.result_summary import TARGETS

ORDER = ('CMD', 'DDD', 'GARDNER', 'WOODS', 'CPZ_FGA', 'MED', 'ED95', 'CMD_DIRECT', 'CMD_INDIRECT')
NAME = 'Results'


def patient_overview(rows, audits, checks):
    by_patient = defaultdict(list)
    for item in audits:
        by_patient[str(item.get('patient', ''))].append(item)
    methods = [m for m in ORDER if any(c['method'] == m for c in checks)]
    output = []
    for row in rows:
        notes = []
        for item in by_patient[str(row['patient_id'])]:
            status = item.get('status')
            if is_excluded(item): notes.append(exclusion_label(item['drug']) + '·환산값 0')
            if status == 'unsupported_formulation': notes.append('주사 간격·제형 확인')
            elif status == 'unknown_drug': notes.append('약물명 확인')
            elif status == 'missing_factor': notes.append('환산계수 없음')
            elif status in ('review', 'missing_unit'): notes.append('용량·단위 확인')
            if item.get('unit_assumed'): notes.append('mg 단위 가정')
            if item.get('frequency') == 'ASSUMED_QD': notes.append('1일 1회 가정')
            if item.get('route') == 'injection' and status == 'converted': notes.append('주사제 환산 가정 확인')
            if item.get('match_type') == 'user_confirmed_alias': notes.append('확인된 약어 적용')
        if 'ID 누락' in row.get('확인할 내용', ''): notes.append('환자 ID 확인')
        if not by_patient[str(row['patient_id'])]: notes.append('약물 입력 확인')
        values = {m:row.get(f'{m} ({TARGETS[m]} mg/day)') for m in methods}
        output.append(dict(id=str(row['patient_id']), values=values, notes=' · '.join(dict.fromkeys(notes))))
    return methods, output


def dated_overview(results):
    grouped = {}
    methods = [m for m in ORDER if any(r['method'] == m for r in results)]
    for result in results:
        key = (result['patient_id'], result['reference_date'])
        row = grouped.setdefault(key, dict(id=key[0], date=key[1], values={}, notes=''))
        row['values'][result['method']] = result['equivalent_mg']
        if result['status'] == 'review': row['notes'] = '상세 시트에서 미환산 사유 확인'
        elif result['status'] == 'non_target': row['notes'] = '병용약물 제외·항정신병약 환산값 0'
        elif result['status'] == 'no_record': row['notes'] = '해당 날짜의 환산 대상 처방 없음'
    return methods, list(grouped.values())


def layout(methods, rows, dated=False, date_label='처방일'):
    methods = [m for m in ORDER if m in methods]
    headers = ['환자 ID'] + ([date_label] if dated else []) + ['환산 상태'] + [m.replace('CMD_', 'CMD\n') + '\n' + TARGETS[m] + ' mg/day' for m in methods] + ['확인할 사항']
    cells = [(0, c, h, 'header') for c, h in enumerate(headers)]
    for i, row in enumerate(rows, 1):
        values = [row['values'].get(m) for m in methods]
        count = sum(v is not None for v in values)
        status = '산출 완료' if count == len(methods) and count else '일부 산출' if count else '계산 보류'
        prefix = [row['id']] + ([row['date']] if dated else []) + [status]
        note = row['notes'] or ('방법별 지원 범위 차이' if count and count < len(methods) else '')
        for c, v in enumerate(prefix): cells.append((i, c, v, 'text'))
        for c, v in enumerate(values, len(prefix)):
            cells.append((i, c, '—' if v is None else v, 'text' if v is None else 'number'))
        cells.append((i, len(headers)-1, note, 'text'))
    return headers, cells, [18] + ([14] if dated else []) + [14] + [16] * len(methods) + [48]


def add_openpyxl(wb, methods, rows):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    from services.excel_tables import text_height
    ws = wb.create_sheet(NAME, 0)
    headers, cells, widths = layout(methods, rows)
    for r, c, value, kind in cells:
        cell = ws.cell(r+1, c+1, value)
        if isinstance(value, str): cell.data_type = 's'
        cell.font = Font(name='맑은 고딕', size=11, bold=kind == 'header', color='222222')
        if kind == 'header': cell.fill = PatternFill('solid', fgColor='EEEEEE')
        cell.alignment = Alignment(vertical='top', wrap_text=True, horizontal='right' if kind == 'number' else 'left')
        if kind == 'number': cell.number_format = '0.00'
        ws.row_dimensions[r+1].height = max(ws.row_dimensions[r+1].height or 20, text_height(value, widths[c]))
    for c, width in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = width
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{len(rows)+1}'
    ws.freeze_panes = None
    ws.sheet_view.showGridLines = True
    ws.sheet_view.zoomScale = 100
    wb.active = 0


def add_xlsxwriter(wb, methods, rows, date_label='처방일', dated=True):
    from services.excel_tables import write_table
    headers, cells, widths = layout(methods, rows, dated=dated, date_label=date_label)
    # Row dictionaries preserve numeric values and literal IDs in the shared writer.
    def values():
        current, row = 1, {}
        for r, c, value, kind in cells:
            if r == 0: continue
            if r != current:
                yield row
                current, row = r, {}
            row[headers[c]] = value
        if row: yield row
    return write_table(wb, NAME, headers, values(), widths=widths)
